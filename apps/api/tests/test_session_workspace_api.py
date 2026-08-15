from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlmodel import Session, select

from app.db import engine
from app.main import app
from app.models import AiModelConfig, LessonPlan, OutlineRow, SessionMaterial
from app.services.secret_store import encrypt_secret
from app.services.session_material_ai import SessionMaterialValidationError
from app.services.session_material_generator import generate_session_material


def _auth_headers(client: TestClient, employee_no: str = "admin", password: str = "Admin@2026!"):
    response = client.post("/auth/login", json={"employee_no": employee_no, "password": password})
    assert response.status_code == 200
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


@pytest.fixture(autouse=True)
def mock_ai_material_generation(monkeypatch):
    def fake_generate(_config, _task, outline, lesson, payload):
        return generate_session_material(
            outline,
            lesson,
            payload.material_type,
            payload.difficulty,
            payload.estimated_minutes,
            payload.question_count,
        )

    monkeypatch.setattr("app.routes.tasks.generate_ai_session_material", fake_generate, raising=False)


@pytest.fixture
def prepared_course():
    with TestClient(app) as client:
        client.headers.update(_auth_headers(client))
        response = client.post(
            "/tasks",
            json={
                "term": "2026-2027 第一学期",
                "major": "数字媒体艺术设计",
                "class_name": "数媒艺术 2501",
                "course_name": "人工智能与创意设计",
                "teacher_name": "张老师",
                "location": "智慧教室",
                "total_hours": 32,
                "hours_per_session": 4,
            },
        )
        assert response.status_code == 200
        task_id = response.json()["id"]

        with Session(engine) as session:
            session.add(
                AiModelConfig(
                    base_url="https://model.example/v1",
                    model_name="lesson-model",
                    encrypted_api_key=encrypt_secret("sk-private"),
                    enabled=True,
                    connection_status="connected",
                )
            )
            outline = OutlineRow(
                task_id=task_id,
                session_no=1,
                date_text="2026-09-07",
                week_no=1,
                weekday="周一",
                periods="1-4",
                topic="AIGC 与创意设计导入",
                teaching_content="理解 AIGC 基础并完成案例拆解",
                post_task="提交案例分析记录",
                course_goal_codes="M1",
                ability_codes="1-3-4",
            )
            session.add(outline)
            session.commit()
            session.refresh(outline)
            assert outline.id is not None
            lesson = LessonPlan(
                task_id=task_id,
                outline_row_id=outline.id,
                session_no=1,
                title="第 1 次课：AIGC 与创意设计导入",
                teaching_goals="理解 AIGC 基础",
                key_points="案例拆解",
                difficult_points="将原理用于设计任务",
                teaching_process="课前阅读，课中拆解，课后总结。",
                homework="提交调研记录",
                reflection="",
                course_goal_codes="M1",
                ability_codes="1-3-4",
            )
            session.add(lesson)
            session.commit()
            session.refresh(lesson)
            assert lesson.id is not None
            outline_id = outline.id
            lesson_id = lesson.id

        yield client, task_id, outline_id, lesson_id


def test_session_workspace_returns_outline_lesson_and_materials(prepared_course):
    client, task_id, outline_row_id, _ = prepared_course

    response = client.get(f"/tasks/{task_id}/sessions/{outline_row_id}")

    assert response.status_code == 200
    assert response.json()["outline"]["id"] == outline_row_id
    assert response.json()["lesson"]["outline_row_id"] == outline_row_id
    assert response.json()["materials"] == []


def test_generating_material_preserves_existing_materials(prepared_course):
    client, task_id, outline_row_id, _ = prepared_course
    endpoint = f"/tasks/{task_id}/sessions/{outline_row_id}/materials/generate"

    first = client.post(
        endpoint,
        json={
            "material_type": "assignment",
            "difficulty": "medium",
            "estimated_minutes": 40,
            "question_count": 5,
        },
    )
    second = client.post(
        endpoint,
        json={
            "material_type": "test",
            "difficulty": "basic",
            "estimated_minutes": 20,
            "question_count": 3,
        },
    )

    assert first.status_code == 200
    assert second.status_code == 200
    workspace = client.get(f"/tasks/{task_id}/sessions/{outline_row_id}").json()
    assert [item["material_type"] for item in workspace["materials"]] == ["assignment", "test"]
    assert workspace["materials"][0]["course_goal_codes"] == "M1"
    assert workspace["materials"][0]["ability_codes"] == "1-3-4"


def test_material_generation_requires_an_enabled_model(prepared_course):
    client, task_id, outline_row_id, _ = prepared_course
    with Session(engine) as session:
        for config in session.exec(select(AiModelConfig)).all():
            session.delete(config)
        session.commit()

    response = client.post(
        f"/tasks/{task_id}/sessions/{outline_row_id}/materials/generate",
        json={"material_type": "assignment", "difficulty": "medium", "estimated_minutes": 40},
    )

    assert response.status_code == 409


def test_ai_material_is_marked_and_validation_failure_does_not_save(prepared_course, monkeypatch):
    client, task_id, outline_row_id, _ = prepared_course
    endpoint = f"/tasks/{task_id}/sessions/{outline_row_id}/materials/generate"

    created = client.post(
        endpoint,
        json={"material_type": "assignment", "difficulty": "medium", "estimated_minutes": 40},
    )
    assert created.status_code == 200
    assert created.json()["generation_method"] == "ai"

    def reject(*_args, **_kwargs):
        raise SessionMaterialValidationError("测试题量必须为 3 题")

    monkeypatch.setattr("app.routes.tasks.generate_ai_session_material", reject)
    failed = client.post(
        endpoint,
        json={"material_type": "test", "difficulty": "medium", "estimated_minutes": 20, "question_count": 3},
    )

    assert failed.status_code == 422
    with Session(engine) as session:
        materials = session.exec(select(SessionMaterial).where(SessionMaterial.task_id == task_id)).all()
    assert len(materials) == 1


def test_updates_and_deletes_material(prepared_course):
    client, task_id, outline_row_id, _ = prepared_course
    generated = client.post(
        f"/tasks/{task_id}/sessions/{outline_row_id}/materials/generate",
        json={"material_type": "assignment", "difficulty": "medium", "estimated_minutes": 40},
    ).json()

    updated = client.put(
        f"/tasks/{task_id}/materials/{generated['id']}",
        json={
            "title": "修改后的实践作业",
            "content": "完成一份 AI 创意设计方案。",
            "reference_answer": "方案需要包含目标、流程和成果。",
            "grading_criteria": "目标 30 分；流程 40 分；成果 30 分。",
            "difficulty": "advanced",
            "estimated_minutes": 60,
        },
    )

    assert updated.status_code == 200
    assert updated.json()["title"] == "修改后的实践作业"
    assert updated.json()["difficulty"] == "advanced"
    deleted = client.delete(f"/tasks/{task_id}/materials/{generated['id']}")
    assert deleted.status_code == 204
    workspace = client.get(f"/tasks/{task_id}/sessions/{outline_row_id}").json()
    assert workspace["materials"] == []


def test_teacher_cannot_access_another_users_session(prepared_course):
    client, task_id, outline_row_id, _ = prepared_course
    employee_no = f"T{task_id:04d}"
    created = client.post(
        "/admin/users",
        json={
            "employee_no": employee_no,
            "name": "其他教师",
            "password": "Teacher@2026!",
            "role": "teacher",
            "major_ids": [],
            "is_active": True,
        },
    )
    assert created.status_code == 200
    client.headers.update(_auth_headers(client, employee_no, "Teacher@2026!"))

    response = client.get(f"/tasks/{task_id}/sessions/{outline_row_id}")

    assert response.status_code == 404


def test_rejects_outline_row_from_another_course(prepared_course):
    client, _, outline_row_id, _ = prepared_course
    response = client.post(
        "/tasks",
        json={
            "term": "2026-2027 第一学期",
            "major": "数字媒体艺术设计",
            "class_name": "数媒艺术 2502",
            "course_name": "另一门课程",
            "teacher_name": "张老师",
            "location": "智慧教室",
            "total_hours": 16,
            "hours_per_session": 4,
        },
    )
    other_task_id = response.json()["id"]

    result = client.get(f"/tasks/{other_task_id}/sessions/{outline_row_id}")

    assert result.status_code == 404


def test_generates_outline_only_material_when_lesson_is_missing(prepared_course):
    client, task_id, outline_row_id, lesson_id = prepared_course
    with Session(engine) as session:
        lesson = session.get(LessonPlan, lesson_id)
        session.delete(lesson)
        session.commit()

    response = client.post(
        f"/tasks/{task_id}/sessions/{outline_row_id}/materials/generate",
        json={"material_type": "test", "difficulty": "basic", "estimated_minutes": 20, "question_count": 3},
    )

    assert response.status_code == 200
    assert response.json()["source_status"] == "outline_only"


def test_exports_one_material_as_docx(prepared_course):
    client, task_id, outline_row_id, _ = prepared_course
    generated = client.post(
        f"/tasks/{task_id}/sessions/{outline_row_id}/materials/generate",
        json={"material_type": "assignment", "difficulty": "medium", "estimated_minutes": 40},
    ).json()

    response = client.post(f"/tasks/{task_id}/materials/{generated['id']}/export")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert "filename*=UTF-8''" in response.headers["content-disposition"]
    document = Document(BytesIO(response.content))
    text_parts = [paragraph.text for paragraph in document.paragraphs]
    text_parts.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    text = "\n".join(text_parts)
    assert "人工智能与创意设计" in text
    assert "第 1 次课" in text
    assert "参考答案" in text
    assert "评分标准" in text
    assert "1-3-4" in text
