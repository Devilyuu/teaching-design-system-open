from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import RGBColor

from app.services.ai_outline_sections import LearningUnit, OutlineSections
from app.services.course_standard_assessment import AssessmentRow, CourseAssessments
from app.services.course_standard_resources import CourseResources
from app.services.docx_exporter import fill_outline_docx
from app.services.outline_template_filler import (
    OutlineTemplateError,
    template_marks_generated_sections,
)


SCHOOL_TEMPLATE = Path(__file__).resolve().parents[3] / "templates" / "课程实施大纲模板.docx"
MARK = "C00000"


def _sections(**overrides) -> OutlineSections:
    defaults = dict(
        course_summary="本课程带学生走完 AIGC 时代的界面设计工作流。",
        teaching_strategy="线上线下混合式+翻转课堂",
        prerequisites="图像处理技术、UI 界面设计",
        learning_outcomes=["熟练使用即时设计完成界面设计。", "掌握 AI 赋能的视觉设计流程。"],
        learning_units=[
            LearningUnit(
                name="一、AI辅助逻辑建模",
                teaching_content="利用大模型辅助生成调研问卷",
                teaching_methods="讲授、讨论",
                key_points="重点：提升调研广度\n难点：撰写有效 Prompt",
                hours=12,
            ),
            LearningUnit(
                name="二、智能原型制作",
                teaching_content="制作低保真原型图",
                teaching_methods="分组合作学习",
                key_points="重点：维护云端组件库\n难点：框架转化为逻辑",
                hours=68,
            ),
        ],
        study_advice="跟着做、想着做、变着做，三步强化训练。",
        academic_integrity="不得抄袭他人作品。",
        attendance="履行请假手续，缺勤后自主学习。",
        classroom_discipline="上课起立问好，迟到从前门进入。",
        assignment_requirements="作业不允许迟交。",
    )
    defaults.update(overrides)
    return OutlineSections(**defaults)


def _resources() -> CourseResources:
    return CourseResources(
        textbooks=["移动UI界面设计（微课版），张晓景，人民邮电出版社（2018）"],
        references=["APP草图+流程图+交互原型设计教程，刘源（2020）", "PS APP UI设计从零开始学，贾浩梅（2022）"],
        online=["https://www.zcool.com.cn/"],
    )


def _assessments() -> CourseAssessments:
    """Shaped like 四、课程评价 in both real standards: 过程考核 merged down."""
    plain = ((0,), (1,), (2,), (3,))
    return CourseAssessments(
        rows=(
            AssessmentRow(("过程考核\n（平时成绩）", "出勤", "按迟到早退请假综合评定。", "10%"), plain, ()),
            AssessmentRow(("过程考核\n（平时成绩）", "单元项目", "按单元项目完成情况评定。", "40%"), plain, (0,)),
            AssessmentRow(("结果考核\n(期末成绩)", "期末项目", "由考核小组评定作品汇报。", "50%"), plain, ()),
            AssessmentRow(("合计", "合计", "合计", "100%"), ((0, 1, 2), (3,)), ()),
        )
    )


def _schedule_row() -> dict:
    return {"date_text": "3.2", "week_no": 1, "periods": "1-4", "teaching_content": "市场调研"}


def _export(tmp_path, template=SCHOOL_TEMPLATE, **kwargs):
    output = tmp_path / "outline.docx"
    fill_outline_docx(
        template,
        output,
        {"课程名称": "移动终端APP设计", "任课教师": "张明"},
        [_schedule_row()],
        **kwargs,
    )
    return Document(output)


def _paragraphs(document) -> list[str]:
    return [paragraph.text.strip() for paragraph in document.paragraphs]


def test_the_school_template_comes_back_with_every_section_written(tmp_path):
    """The whole outline is generated now, not just its schedule."""
    rendered = _export(tmp_path, sections=_sections(), resources=_resources(), assessments=_assessments())

    text = "\n".join(_paragraphs(rendered))
    assert "课程简介：本课程带学生走完 AIGC 时代的界面设计工作流。" in text
    assert "教学策略：线上线下混合式+翻转课堂" in text
    assert "先修要求：图像处理技术、UI 界面设计" in text
    assert "跟着做、想着做、变着做，三步强化训练。" in text
    assert "不得抄袭他人作品。" in text
    assert "履行请假手续，缺勤后自主学习。" in text
    assert "上课起立问好，迟到从前门进入。" in text
    assert "作业不允许迟交。" in text
    # Nothing may still be asking to be generated.
    assert "AI 生成" not in text
    assert "AI 生成" not in "\n".join(
        cell.text for table in rendered.tables for row in table.rows for cell in row.cells
    )


def test_the_red_marking_never_reaches_the_finished_document(tmp_path):
    """The colour is the school's instruction to us, not part of its document.

    It marks whole lines, lead-ins included, and paragraph marks besides -- a
    red paragraph mark turns the teacher's own next sentence red.
    """
    assert MARK in Document(SCHOOL_TEMPLATE).element.body.xml

    rendered = _export(tmp_path, sections=_sections(), resources=_resources(), assessments=_assessments())

    assert MARK not in rendered.element.body.xml


def test_a_list_repeats_the_line_the_template_says_to_copy(tmp_path):
    """「按实际条数复制本行」 -- and the label heads the list only once."""
    rendered = _export(tmp_path, sections=_sections(), resources=_resources(), assessments=_assessments())

    text = _paragraphs(rendered)
    assert "熟练使用即时设计完成界面设计。" in text
    assert "掌握 AI 赋能的视觉设计流程。" in text
    # 「教材：」「教辅：」 head their lists from a line of their own, and the
    # entries are cited 「[1] 」「[2] 」 the way the school's finished outlines do.
    textbook = text.index("教材：")
    assert text[textbook + 1] == "[1] 移动UI界面设计（微课版），张晓景，人民邮电出版社（2018）"
    reference = text.index("教辅：")
    assert text[reference + 1 : reference + 3] == [
        "[1] APP草图+流程图+交互原型设计教程，刘源（2020）",
        "[2] PS APP UI设计从零开始学，贾浩梅（2022）",
    ]
    assert not any(line.startswith("教辅：APP") for line in text)


def test_an_entry_the_standard_already_numbers_is_not_numbered_again(tmp_path):
    resources = CourseResources(textbooks=["[1] 自编讲义"], references=["1．参考书甲", "2．参考书乙"])

    rendered = _export(tmp_path, sections=_sections(), resources=resources, assessments=_assessments())

    text = _paragraphs(rendered)
    assert "[1] 自编讲义" in text
    assert "1．参考书甲" in text
    assert not any(line.startswith("[1] 1．") or line.startswith("[1] [1]") for line in text)


def test_each_generated_cell_keeps_the_format_its_own_marker_wore(tmp_path):
    """教学单元 is a size up from the columns beside it in the school's outlines."""
    rendered = _export(tmp_path, sections=_sections(), resources=_resources(), assessments=_assessments())

    content = next(table for table in rendered.tables if "教学单元" in table.rows[0].cells[0].text)
    unit_run = content.rows[1].cells[0].paragraphs[0].runs[0]
    detail_run = content.rows[1].cells[1].paragraphs[0].runs[0]
    assert unit_run.text == "一、AI辅助逻辑建模"
    assert unit_run._r.rPr.find(qn("w:sz")).get(qn("w:val")) == "21"
    assert detail_run._r.rPr.find(qn("w:sz")).get(qn("w:val")) == "18"
    assert unit_run._r.rPr.find(qn("w:color")) is None


def test_a_subsection_the_standard_says_nothing_about_asks_the_teacher(tmp_path):
    """Padding it out with a neighbour's entry is the same defect as inventing it."""
    rendered = _export(tmp_path, sections=_sections(), resources=_resources(), assessments=_assessments())

    assert "课程标准中未提供，请教师补充" in _paragraphs(rendered)


def test_the_learning_content_table_grows_one_row_per_unit(tmp_path):
    rendered = _export(tmp_path, sections=_sections(), resources=_resources(), assessments=_assessments())

    content = next(table for table in rendered.tables if "教学单元" in table.rows[0].cells[0].text)
    assert len(content.rows) == 3
    assert content.cell(1, 0).text == "一、AI辅助逻辑建模"
    assert content.cell(1, 3).text == "重点：提升调研广度\n难点：撰写有效 Prompt"
    # The hours are the course standard's own, and they are what gets checked
    # against the course total.
    assert [content.cell(row, 4).text for row in (1, 2)] == ["12", "68"]


def _tc_shape(row) -> list[tuple[str | None, str | None]]:
    """(gridSpan, vMerge) per w:tc, which is how Word records the merges."""
    shape = []
    for cell in row._tr.findall(qn("w:tc")):
        properties = cell.find(qn("w:tcPr"))
        span = None if properties is None else properties.find(qn("w:gridSpan"))
        merge = None if properties is None else properties.find(qn("w:vMerge"))
        shape.append((
            None if span is None else span.get(qn("w:val")),
            None if merge is None else (merge.get(qn("w:val")) or "continue"),
        ))
    return shape


def test_the_assessment_table_is_copied_from_the_standard(tmp_path):
    """The school decided this split of the marks, and the 教务 record checks it."""
    rendered = _export(tmp_path, sections=_sections(), resources=_resources(), assessments=_assessments())

    assessment = next(table for table in rendered.tables if "考核项目" in table.rows[0].cells[0].text)
    # The header runs 考核项目 across the first two grid columns, so cell 1 of the
    # header is still 考核项目; the rows below fill both columns of that span.
    assert [row.cells[1].text for row in assessment.rows[1:]] == ["出勤", "单元项目", "期末项目", "合计"]
    assert [row.cells[0].text for row in assessment.rows[1:3]] == ["过程考核\n（平时成绩）"] * 2
    assert assessment.cell(1, 2).text == "按迟到早退请假综合评定。"
    assert [assessment.cell(row, 3).text for row in (1, 2, 3, 4)] == ["10%", "40%", "50%", "100%"]
    # 过程考核 runs down two rows in the standard, and 合计 across three columns.
    assert _tc_shape(assessment.rows[1])[0] == (None, "restart")
    assert _tc_shape(assessment.rows[2])[0] == (None, "continue")
    assert _tc_shape(assessment.rows[3])[0] == (None, None)
    assert _tc_shape(assessment.rows[4]) == [("3", None), (None, None)]


def test_the_assessment_table_says_so_when_the_standard_is_silent(tmp_path):
    """Inventing a split of the marks would be the same defect as inventing a textbook."""
    rendered = _export(tmp_path, sections=_sections(), resources=_resources())

    assessment = next(table for table in rendered.tables if "考核项目" in table.rows[0].cells[0].text)
    assert len(assessment.rows) == 2
    assert assessment.cell(1, 0).text == "课程标准中未提供，请教师补充"
    assert _tc_shape(assessment.rows[1]) == [("4", None)]


def test_the_schedule_is_still_written_alongside_the_body(tmp_path):
    rendered = _export(tmp_path, sections=_sections(), resources=_resources(), assessments=_assessments())

    schedule = next(table for table in rendered.tables if "日期" in table.rows[0].cells[0].text)
    assert schedule.cell(1, 0).text == "3.2"
    assert schedule.cell(1, 3).text == "市场调研"


def test_the_export_stops_when_the_body_was_never_generated(tmp_path):
    """Filling only the schedule would ship 「【AI 生成】」 as the course's text."""
    with pytest.raises(OutlineTemplateError, match="还没有生成大纲正文"):
        _export(tmp_path)

    assert not (tmp_path / "outline.docx").exists()


def _marked_template(path, extra_paragraph: str | None = None, total_row: bool = False):
    document = Document()
    document.add_paragraph("二、课程介绍")
    _mark(document.add_paragraph(), "课程简介：【AI 生成】")
    document.add_paragraph("六、考核方式与评价标准")
    table = document.add_table(rows=2, cols=3)
    for index, header in enumerate(["考核项目", "考核方式", "比例"]):
        table.cell(0, index).text = header
    for index in range(3):
        _mark(table.cell(1, index).paragraphs[0], "【AI 生成】")
    if total_row:
        row = table.add_row()
        row.cells[0].text = "合计"
        row.cells[2].text = "100%"
    if extra_paragraph is not None:
        document.add_paragraph(extra_paragraph)
        _mark(document.add_paragraph(), "【AI 生成】")
    document.save(path)


def _mark(paragraph, text: str) -> None:
    run = paragraph.add_run(text)
    run.font.color.rgb = RGBColor.from_string(MARK)


def test_a_marker_nobody_wired_up_stops_the_export(tmp_path):
    """The gate that caught 60 real gaps in the lesson template, for the outline."""
    template = tmp_path / "extra-marker.docx"
    _marked_template(template, extra_paragraph="九、企业专家寄语")

    with pytest.raises(OutlineTemplateError, match="仍是模板占位符"):
        _export(tmp_path, template=template, sections=_sections(), resources=_resources())


def test_a_narrower_template_is_refused_rather_than_dropping_the_percentages(tmp_path):
    """比例 is the standard's last column; writing what fits still looks finished."""
    template = tmp_path / "narrow.docx"
    _marked_template(template)

    with pytest.raises(OutlineTemplateError, match="丢掉最后一列"):
        _export(tmp_path, template=template, sections=_sections(), assessments=_assessments())


def test_a_total_row_the_template_drew_itself_survives(tmp_path):
    """Only marked rows are cloned or dropped, so an unmarked 合计 row stays put."""
    template = tmp_path / "with-total.docx"
    _marked_template(template, total_row=True)

    rendered = _export(tmp_path, template=template, sections=_sections(), resources=_resources())

    assessment = rendered.tables[0]
    assert [row.cells[0].text for row in assessment.rows] == [
        "考核项目",
        "课程标准中未提供，请教师补充",
        "合计",
    ]
    assert assessment.rows[-1].cells[2].text == "100%"


def test_a_template_without_markers_is_left_to_its_own_words(tmp_path):
    """A teacher's own finished outline still exports as it did before."""
    template = tmp_path / "plain.docx"
    document = Document()
    document.add_paragraph("二、课程介绍")
    document.add_paragraph("课程简介：这门课由我自己写。")
    table = document.add_table(rows=2, cols=4)
    for index, header in enumerate(["日期", "周次", "节次", "教学内容"]):
        table.cell(0, index).text = header
    document.save(template)

    assert template_marks_generated_sections(Document(template)) is False

    rendered = _export(tmp_path, template=template)

    assert "课程简介：这门课由我自己写。" in _paragraphs(rendered)
    assert rendered.tables[0].cell(1, 0).text == "3.2"
