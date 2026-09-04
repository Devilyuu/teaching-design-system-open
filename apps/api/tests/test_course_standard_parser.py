from pathlib import Path

from docx import Document

from app.services.course_standard_parser import parse_course_standard


def make_course_standard_docx(path: Path) -> None:
    doc = Document()
    table = doc.add_table(rows=5, cols=3)
    table.cell(0, 0).text = "编号"
    table.cell(0, 1).text = "课程教学目标"
    table.cell(0, 2).text = "对应的知识能力素养集"
    table.cell(1, 0).text = "M1"
    table.cell(1, 1).text = "理解AIGC基础与创意设计流程"
    table.cell(1, 2).text = "1-3-4, 1-3-5"
    table.cell(2, 0).text = "M2"
    table.cell(2, 1).text = "完成AI辅助文创项目"
    table.cell(2, 2).text = "2-3-4 2-3-5"
    table.cell(3, 0).text = "M3"
    table.cell(3, 1).text = "形成审美判断与表达能力"
    table.cell(3, 2).text = "3-2-1, 3-3-3"
    table.cell(4, 0).text = "M4"
    table.cell(4, 1).text = "塑造求实创新精神"
    table.cell(4, 2).text = "3-5-1"
    doc.save(path)


def make_project_course_standard_docx(path: Path) -> None:
    doc = Document()
    goal_table = doc.add_table(rows=3, cols=3)
    goal_table.cell(0, 0).text = "编号"
    goal_table.cell(0, 1).text = "课程教学目标"
    goal_table.cell(0, 2).text = "对应的知识能力素养集"
    goal_table.cell(1, 0).text = "M1"
    goal_table.cell(1, 1).text = "理解 AIGC 基础"
    goal_table.cell(1, 2).text = "1-3-4"
    goal_table.cell(2, 0).text = "M2"
    goal_table.cell(2, 1).text = "完成 AI 创意项目"
    goal_table.cell(2, 2).text = "2-3-4"

    table = doc.add_table(rows=5, cols=7)
    headers = ["项目名称", "教学内容", "教学方法建议", "教学目标", "知识能力素养集", "知识能力素养集", "参考课时"]
    for index, header in enumerate(headers):
        table.cell(0, index).text = header
    rows = [
        ["项目一：工具探索", "了解生成式 AI 并完成首次图像生成。", "讲授 案例分析", "M1 M2", "1-3-4 2-3-4", "1-3-4 2-3-4", "8/4"],
        ["项目一：工具探索", "项目描述：本项目为入门阶段。", "项目描述：本项目为入门阶段。", "项目描述：本项目为入门阶段。", "项目描述：本项目为入门阶段。", "项目描述：本项目为入门阶段。", "项目描述：本项目为入门阶段。"],
        ["项目二：视觉开发", "完成剧本、分镜和视觉概念图。", "工作坊 任务驱动", "M2", "2-3-4", "2-3-4", "24/12"],
        ["合计", "合计", "合计", "合计", "合计", "合计", "32/16"],
    ]
    for row_index, values in enumerate(rows, start=1):
        for column_index, value in enumerate(values):
            table.cell(row_index, column_index).text = value
    doc.save(path)


def test_parse_course_goals_from_standard_docx(tmp_path):
    path = tmp_path / "standard.docx"
    make_course_standard_docx(path)

    result = parse_course_standard(path)

    assert [goal.code for goal in result.goals] == ["M1", "M2", "M3", "M4"]
    assert result.goals[0].ability_codes == ["1-3-4", "1-3-5"]
    assert result.goals[1].ability_codes == ["2-3-4", "2-3-5"]
    assert result.goals[3].description == "塑造求实创新精神"


def test_parse_teaching_projects_and_merge_description_rows(tmp_path):
    path = tmp_path / "project-standard.docx"
    make_project_course_standard_docx(path)

    result = parse_course_standard(path)

    assert len(result.projects) == 2
    assert result.projects[0].name == "项目一：工具探索"
    assert result.projects[0].reference_hours == 8
    assert result.projects[0].practice_hours == 4
    assert result.projects[0].course_goal_codes == ["M1", "M2"]
    assert result.projects[0].ability_codes == ["1-3-4", "2-3-4"]
    assert result.projects[0].description == "本项目为入门阶段。"
    assert result.projects[1].reference_hours == 24


def test_parse_projects_when_hours_column_is_a_bare_total(tmp_path):
    # Some schools write "12" instead of "8/4"; the outline only needs the total.
    path = tmp_path / "total-hours-standard.docx"
    doc = Document()
    table = doc.add_table(rows=4, cols=7)
    headers = ["项目名称", "教学内容", "教学方法建议", "教学目标", "知识能力素养集", "知识能力素养集", "参考课时"]
    for index, header in enumerate(headers):
        table.cell(0, index).text = header
    rows = [
        ["项目一：节日海报", "掌握动态海报表现方法", "讲授 练习", "M1", "1-1-3", "1-1-3", "12"],
        ["项目二：综合应用", "完成设计稿制作", "讲授 分组合作", "M3 M4", "1-4-5 3-1-4", "1-4-5 3-1-4", "36 学时"],
        ["合计", "合计", "合计", "合计", "合计", "合计", "48"],
    ]
    for row_index, values in enumerate(rows, start=1):
        for column_index, value in enumerate(values):
            table.cell(row_index, column_index).text = value
    doc.save(path)

    result = parse_course_standard(path)

    assert [project.name for project in result.projects] == ["项目一：节日海报", "项目二：综合应用"]
    assert [project.reference_hours for project in result.projects] == [12, 36]
    assert [project.practice_hours for project in result.projects] == [0, 0]
    assert result.projects[1].course_goal_codes == ["M3", "M4"]
