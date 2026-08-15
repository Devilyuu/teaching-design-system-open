from types import SimpleNamespace

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.services.lesson_template_filler import (
    LessonTemplateError,
    fill_structured_lesson_template,
    find_lesson_template_blocks,
    find_unfilled_generated_cells,
    read_process_layout,
)


def _merge_text(table, row: int, start: int, end: int, text: str) -> None:
    cell = table.cell(row, start)
    if end > start:
        cell = cell.merge(table.cell(row, end))
    cell.text = text


def _add_lesson_block(doc: Document, title: str) -> None:
    info = doc.add_table(rows=8, cols=10)
    _merge_text(info, 0, 0, 2, "课程名称")
    _merge_text(info, 0, 3, 5, "模板课程")
    info.cell(0, 6).text = "任课教师"
    _merge_text(info, 0, 7, 9, "模板教师")
    _merge_text(info, 2, 0, 2, "本次课标题")
    _merge_text(info, 2, 3, 5, title)
    info.cell(2, 6).text = "授课学时"
    _merge_text(info, 2, 7, 9, "4")
    info.cell(3, 0).text = "授课班级"
    _merge_text(info, 3, 1, 3, "模板班级")
    info.cell(3, 4).text = "上课时间"
    _merge_text(info, 3, 5, 6, "第 1 周 周一 第 1-4 节")
    info.cell(3, 7).text = "上课地点"
    _merge_text(info, 3, 8, 9, "模板教室")
    _merge_text(info, 6, 0, 6, "本次课教学目标")
    _merge_text(info, 6, 7, 8, "课程教学目标")
    info.cell(6, 9).text = "能力指标代码"
    _merge_text(info, 7, 0, 6, "模板教学目标")
    _merge_text(info, 7, 7, 8, "M1")
    info.cell(7, 9).text = "1-1-1"

    process = doc.add_table(rows=7, cols=6)
    headers = ["教学环节（时长）", "时长", "教学内容", "教师活动", "学生活动", "对应教学目标"]
    for index, header in enumerate(headers):
        process.cell(0, index).text = header
    phases = ["课前", "课中", "课中", "课中", "课中", "课后"]
    for row_index, phase in enumerate(phases, start=1):
        process.cell(row_index, 0).text = phase
        process.cell(row_index, 1).text = ""
        process.cell(row_index, 2).text = "模板教学内容"
        process.cell(row_index, 3).text = "模板教师活动"
        process.cell(row_index, 4).text = "模板学生活动"
        process.cell(row_index, 5).text = "M1"

    reflection = doc.add_table(rows=1, cols=1)
    reflection.cell(0, 0).text = "教学反思"


def _structure_signature(document: Document) -> list[tuple[int, int, tuple[str, ...]]]:
    signature = []
    for table in document.tables:
        merge_xml = tuple(
            element.xml
            for row in table.rows
            for cell in row.cells
            for element in cell._tc.tcPr
            if element.tag.endswith("gridSpan") or element.tag.endswith("vMerge")
        )
        signature.append((len(table.rows), len(table.columns), merge_xml))
    return signature


def _lesson(number: int):
    return SimpleNamespace(
        outline_row_id=number,
        title=f"第 {number} 次课：AI 创意设计 {number}",
        duration_minutes=160,
        teaching_goals=f"完成第 {number} 次课项目任务",
        teaching_preparation=f"教师准备案例 {number}，学生准备设计素材",
        teaching_process="\n\n".join(
            [
                "\n".join(
                    [
                        "导入（20分钟）",
                        f"教师活动：展示案例 {number}",
                        f"学生活动：观察并提出问题 {number}",
                        "学习评价：口头提问",
                    ]
                ),
                "\n".join(
                    [
                        "实训（140分钟）",
                        f"教师活动：示范并巡回指导 {number}",
                        f"学生活动：完成设计作品 {number}",
                        "学习评价：过程检查",
                    ]
                ),
            ]
        ),
        homework=f"提交作品 {number}",
        course_goal_codes=f"M{number}",
        ability_codes=f"{number}-1-1",
    )


def test_finds_repeated_lesson_blocks_without_fixed_table_indexes():
    doc = Document()
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "无关表格"
    _add_lesson_block(doc, "模板课次一")
    _add_lesson_block(doc, "模板课次二")

    blocks = find_lesson_template_blocks(doc)

    assert len(blocks) == 2
    assert blocks[0].info_table_index == 1
    assert blocks[0].process_table_index == 2
    assert blocks[1].info_table_index == 4
    assert blocks[1].process_table_index == 5


def test_fills_each_lesson_in_place_and_preserves_table_structure():
    doc = Document()
    _add_lesson_block(doc, "模板课次一")
    _add_lesson_block(doc, "模板课次二")
    before = _structure_signature(doc)

    fill_structured_lesson_template(
        doc,
        [_lesson(1), _lesson(2)],
        [
            {"week_no": 3, "weekday": "周一", "periods": "1-4", "class_name": "数媒 2501", "location": "A101"},
            {"week_no": 4, "weekday": "周二", "periods": "5-8", "class_name": "数媒 2501", "location": "A102"},
        ],
        {"课程名称": "人工智能与创意设计", "任课教师": "张老师"},
    )

    assert _structure_signature(doc) == before
    first_info, first_process = doc.tables[0], doc.tables[1]
    second_info, second_process = doc.tables[3], doc.tables[4]
    assert first_info.cell(0, 3).text == "人工智能与创意设计"
    assert first_info.cell(0, 7).text == "张老师"
    assert first_info.cell(2, 3).text == "第 1 次课：AI 创意设计 1"
    assert first_info.cell(2, 7).text == "4"
    assert first_info.cell(3, 1).text == "数媒 2501"
    assert first_info.cell(3, 5).text == "第 3 周 周一 第 1-4 节"
    assert first_info.cell(3, 8).text == "A101"
    assert first_info.cell(7, 0).text == "K1：完成第 1 次课项目任务"
    assert first_info.cell(7, 7).text == "M1"
    assert first_info.cell(7, 9).text == "1-1-1"
    assert first_process.cell(1, 2).text == "教学准备"
    assert first_process.cell(1, 3).text == "发布课前学习任务并检查学习资源。"
    assert first_process.cell(1, 4).text == "预习本次课内容，做好上课准备。"
    assert first_process.cell(2, 1).text == "20分钟"
    assert "导入" in first_process.cell(2, 2).text
    assert first_process.cell(2, 3).text == "展示案例 1"
    assert first_process.cell(2, 4).text == "观察并提出问题 1"
    assert first_process.cell(6, 2).text == "课后巩固与提交"
    assert first_process.cell(6, 3).text == "提交作品 1"
    assert first_process.cell(6, 4).text == "提交作品 1"
    # The column names the lesson's own K objectives, not the course-level M codes.
    assert first_process.cell(2, 5).text == "K1"
    assert second_info.cell(2, 3).text == "第 2 次课：AI 创意设计 2"
    assert second_info.cell(3, 5).text == "第 4 周 周二 第 5-8 节"
    assert second_process.cell(3, 3).text == "示范并巡回指导 2"


def test_replaces_all_existing_template_goal_rows():
    doc = Document()
    info = doc.add_table(rows=11, cols=10)
    _merge_text(info, 0, 0, 2, "课程名称")
    _merge_text(info, 0, 3, 5, "模板课程")
    info.cell(0, 6).text = "任课教师"
    _merge_text(info, 0, 7, 9, "模板教师")
    _merge_text(info, 2, 0, 2, "本次课标题")
    _merge_text(info, 2, 3, 5, "模板课次")
    info.cell(2, 6).text = "授课学时"
    _merge_text(info, 2, 7, 9, "4")
    _merge_text(info, 6, 0, 6, "本次课教学目标")
    _merge_text(info, 6, 7, 8, "课程教学目标")
    info.cell(6, 9).text = "能力指标代码"
    for row_index, label in zip((7, 8, 9), ("K1：模板知识目标", "K2：模板能力目标", "K3：模板素养目标")):
        _merge_text(info, row_index, 0, 6, label)
        _merge_text(info, row_index, 7, 8, "M9")
        info.cell(row_index, 9).text = "9-9-9"
    _merge_text(info, 10, 0, 1, "教学重点")
    _merge_text(info, 10, 2, 9, "模板重点")

    process = doc.add_table(rows=7, cols=6)
    for index, header in enumerate(["教学环节", "时长", "教学内容", "教师活动", "学生活动", "对应教学目标"]):
        process.cell(0, index).text = header
    for row_index, phase in enumerate(["课前", "课中", "课中", "课中", "课中", "课后"], start=1):
        process.cell(row_index, 0).text = phase

    fill_structured_lesson_template(doc, [_lesson(1)], [{}], {})

    assert info.cell(7, 0).text == "K1：完成第 1 次课项目任务"
    assert info.cell(7, 7).text == "M1"
    assert info.cell(7, 9).text == "1-1-1"
    assert info.cell(8, 0).text == ""
    assert info.cell(8, 7).text == ""
    assert info.cell(8, 9).text == ""
    assert info.cell(9, 0).text == ""


def test_rejects_a_template_with_no_recognizable_lesson_block():
    unrecognized = Document()
    unrecognized.add_table(rows=1, cols=1).cell(0, 0).text = "普通表格"

    with pytest.raises(LessonTemplateError, match="未识别"):
        fill_structured_lesson_template(unrecognized, [_lesson(1)], [{}], {})


def test_converts_filled_floating_tables_to_inline_flow():
    doc = Document()
    _add_lesson_block(doc, "模板课次")
    for table in doc.tables[:2]:
        positioning = OxmlElement("w:tblpPr")
        positioning.set(qn("w:vertAnchor"), "page")
        positioning.set(qn("w:tblpY"), "2708")
        table._tbl.tblPr.append(positioning)

    fill_structured_lesson_template(doc, [_lesson(1)], [{}], {})

    assert doc.tables[0]._tbl.tblPr.find(qn("w:tblpPr")) is None
    assert doc.tables[1]._tbl.tblPr.find(qn("w:tblpPr")) is None


def test_repeats_a_single_block_template_for_every_session():
    """The school's real template carries one lesson block, not one per session."""
    doc = Document()
    _add_lesson_block(doc, "模板课次")
    lessons = [_lesson(number) for number in (1, 2, 3)]
    contexts = [{"week_no": number, "weekday": "一", "periods": "1-4",
                 "class_name": f"班级{number}", "location": f"教室{number}"} for number in (1, 2, 3)]

    fill_structured_lesson_template(doc, lessons, contexts, {"课程名称": "移动终端APP设计"})

    blocks = find_lesson_template_blocks(doc)
    assert len(blocks) == 3
    titles = [doc.tables[block.info_table_index].cell(2, 3).text for block in blocks]
    assert titles == [lesson.title for lesson in lessons]


def test_repeated_blocks_keep_the_template_structure():
    doc = Document()
    _add_lesson_block(doc, "模板课次")
    reference = _structure_signature(doc)

    fill_structured_lesson_template(
        doc,
        [_lesson(1), _lesson(2)],
        [{}, {}],
        {"课程名称": "移动终端APP设计"},
    )

    produced = _structure_signature(doc)
    assert produced[: len(reference)] == reference
    assert produced[len(reference) :] == reference


def test_reads_the_in_class_segment_layout_the_template_fixes():
    doc = Document()
    _add_lesson_block(doc, "模板课次")
    process = doc.tables[1]
    for row_index, minutes in ((2, "20分钟"), (3, "40分钟"), (4, "60分钟"), (5, "40 分钟")):
        process.cell(row_index, 1).text = minutes

    layout = read_process_layout(doc)

    assert layout == [20, 40, 60, 40]


def test_layout_is_empty_when_the_template_leaves_minutes_blank():
    doc = Document()
    _add_lesson_block(doc, "模板课次")

    assert read_process_layout(doc) == []


def _mark_generated(cell, text: str) -> None:
    from docx.shared import RGBColor

    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)


def test_reports_a_red_marked_cell_the_export_never_filled():
    template = Document()
    _add_lesson_block(template, "模板课次")
    _mark_generated(template.tables[0].cell(7, 0), "模板教学目标")

    produced = Document()
    _add_lesson_block(produced, "模板课次")

    stale = find_unfilled_generated_cells(template, produced)

    assert len(stale) == 1
    assert "第 1 次课" in stale[0]


def test_reports_nothing_once_the_marked_cell_is_replaced():
    template = Document()
    _add_lesson_block(template, "模板课次")
    _mark_generated(template.tables[0].cell(7, 0), "模板教学目标")

    produced = Document()
    _add_lesson_block(produced, "模板课次")
    fill_structured_lesson_template(produced, [_lesson(1)], [{}], {})

    assert find_unfilled_generated_cells(template, produced) == []
