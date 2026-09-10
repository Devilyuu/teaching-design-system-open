from pathlib import Path

from docx import Document

from app.services.talent_plan_parser import parse_talent_plan


def make_talent_plan_docx(path: Path) -> None:
    doc = Document()

    decoy = doc.add_table(rows=2, cols=4)
    decoy.cell(0, 0).text = "类别"
    decoy.cell(0, 1).text = "培养规格代码"
    decoy.cell(0, 2).text = "TOP10"
    decoy.cell(0, 3).text = "其他"
    decoy.cell(1, 1).text = "不是二级代码"
    decoy.cell(1, 2).text = "9-9-9干扰内容"

    target = doc.add_table(rows=7, cols=4)
    target.cell(0, 0).text = "知识、能力与素养指标"
    target.cell(1, 0).text = "类别"
    target.cell(1, 1).text = "培养规格代码"
    target.cell(1, 2).text = "TOP10"
    target.cell(1, 3).text = "其他"

    target.cell(2, 0).text = "知识点"
    target.cell(2, 1).text = "1-1"
    target.cell(2, 2).text = (
        "1-1-1探索基础动画关键帧的制作方法\n"
        "1-1-10总结渲染输出的方法"
    )

    target.cell(3, 0).text = "能力点"
    target.cell(3, 1).text = "2-3"
    target.cell(3, 3).text = "2-3-4熟练运用特效软件完成视频处理"

    target.cell(4, 0).text = "素养点"
    target.cell(4, 1).text = "3-2"
    target.cell(4, 2).text = "3-2-1形成求实创新的职业素养"

    target.cell(5, 0).text = "知识点"
    target.cell(5, 1).text = "1-1-10"
    target.cell(5, 2).text = "8-8-8无效数据行中的内容"

    target.cell(6, 0).text = "知识点"
    target.cell(6, 1).text = "1-1"
    target.cell(6, 2).text = "1-1-1重复代码不应覆盖首次描述"

    doc.save(path)


def test_parse_authoritative_indicator_table_without_fixed_table_index(tmp_path):
    path = tmp_path / "talent-plan.docx"
    make_talent_plan_docx(path)

    result = parse_talent_plan(path)

    assert [item.code for item in result.indicators] == [
        "1-1-1",
        "1-1-10",
        "2-3-4",
        "3-2-1",
    ]
    assert [item.category for item in result.indicators] == [
        "知识点",
        "知识点",
        "能力点",
        "素养点",
    ]
    assert [item.group_code for item in result.indicators] == [
        "1-1",
        "1-1",
        "2-3",
        "3-2",
    ]
    assert result.indicators[0].description == "探索基础动画关键帧的制作方法"
    assert result.indicators[1].description == "总结渲染输出的方法"
    assert result.indicators[2].description == "熟练运用特效软件完成视频处理"
    assert result.indicators[3].description == "形成求实创新的职业素养"


def test_ignore_wrong_group_indicator_before_deduplication(tmp_path):
    path = tmp_path / "wrong-group.docx"
    doc = Document()
    table = doc.add_table(rows=4, cols=4)
    table.cell(0, 1).text = "培养规格代码"
    table.cell(0, 2).text = "TOP10"
    table.cell(0, 3).text = "其他"
    table.cell(1, 0).text = "知识点"
    table.cell(1, 1).text = "1-1"
    table.cell(1, 2).text = "2-3-4错误分组中的描述"
    table.cell(2, 0).text = "能力点"
    table.cell(2, 1).text = "2-3"
    table.cell(2, 2).text = "2-3-4正确分组中的描述"
    table.cell(3, 0).text = "能力点"
    table.cell(3, 1).text = "2-3"
    table.cell(3, 2).text = "2-3-5另一条正确指标"
    doc.save(path)

    result = parse_talent_plan(path)

    assert [
        (item.code, item.category, item.group_code, item.description)
        for item in result.indicators
    ] == [
        ("2-3-4", "能力点", "2-3", "正确分组中的描述"),
        ("2-3-5", "能力点", "2-3", "另一条正确指标"),
    ]


def test_inherit_category_and_group_for_blank_continuation_row(tmp_path):
    path = tmp_path / "continuation.docx"
    doc = Document()
    table = doc.add_table(rows=3, cols=4)
    table.cell(0, 1).text = "培养规格代码"
    table.cell(0, 2).text = "TOP10"
    table.cell(0, 3).text = "其他"
    table.cell(1, 0).text = "知识点"
    table.cell(1, 1).text = "1-2"
    table.cell(1, 2).text = "1-2-1首行指标"
    table.cell(2, 2).text = "1-2-2空白续行指标"
    doc.save(path)

    result = parse_talent_plan(path)

    assert [
        (item.code, item.category, item.group_code, item.description)
        for item in result.indicators
    ] == [
        ("1-2-1", "知识点", "1-2", "首行指标"),
        ("1-2-2", "知识点", "1-2", "空白续行指标"),
    ]


def test_cross_group_code_does_not_split_description(tmp_path):
    path = tmp_path / "description-boundary.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=4)
    table.cell(0, 1).text = "培养规格代码"
    table.cell(0, 2).text = "TOP10"
    table.cell(0, 3).text = "其他"
    table.cell(1, 0).text = "知识点"
    table.cell(1, 1).text = "1-1"
    table.cell(1, 2).text = (
        "1-1-1掌握跨组示例2-3-4并完成原型 "
        "1-1-2理解同组下一指标"
    )
    doc.save(path)

    result = parse_talent_plan(path)

    assert [item.code for item in result.indicators] == ["1-1-1", "1-1-2"]
    assert result.indicators[0].description == "掌握跨组示例2-3-4并完成原型"
    assert result.indicators[1].description == "理解同组下一指标"


def test_invalid_nonempty_group_clears_continuation_context(tmp_path):
    path = tmp_path / "invalid-group-boundary.docx"
    doc = Document()
    table = doc.add_table(rows=4, cols=4)
    table.cell(0, 1).text = "培养规格代码"
    table.cell(0, 2).text = "TOP10"
    table.cell(0, 3).text = "其他"
    table.cell(1, 0).text = "知识点"
    table.cell(1, 1).text = "1-1"
    table.cell(1, 2).text = "1-1-1有效指标"
    table.cell(2, 0).text = "说明"
    table.cell(2, 1).text = "不适用"
    table.cell(3, 2).text = "1-1-2不可越过无效行继承"
    doc.save(path)

    result = parse_talent_plan(path)

    assert [item.code for item in result.indicators] == ["1-1-1"]


def test_group_cell_may_carry_the_group_name(tmp_path):
    """2026 plans write "1-1 艺术设计基础知识" where older ones wrote "1-1"."""
    path = tmp_path / "named-groups.docx"
    doc = Document()
    table = doc.add_table(rows=4, cols=4)
    table.cell(0, 1).text = "培养规格代码"
    table.cell(0, 2).text = "TOP10"
    table.cell(0, 3).text = "其他"
    table.cell(1, 0).text = "知识点"
    table.cell(1, 1).text = "1-1 艺术设计基础知识"
    table.cell(1, 2).text = "1-1-1 理解艺术设计的基本概念 1-1-2 掌握造型基础"
    table.cell(2, 0).text = "知识点"
    table.cell(2, 1).text = "1-10 知识产权、安全与科技伦理知识"
    table.cell(2, 2).text = "1-10-1 掌握著作权的基本知识"
    table.cell(3, 0).text = "能力点"
    table.cell(3, 1).text = "2-1-3 不是分组而是指标编号"
    table.cell(3, 2).text = "2-1-4 不应被收进任何分组"
    doc.save(path)

    result = parse_talent_plan(path)

    assert [item.code for item in result.indicators] == ["1-1-1", "1-1-2", "1-10-1"]
    assert result.indicators[0].group_code == "1-1"
    assert result.indicators[2].group_code == "1-10"
    assert result.indicators[2].description == "掌握著作权的基本知识"
