from types import SimpleNamespace

from docx import Document

from app.services.session_material_exporter import export_session_material_docx


def test_exports_material_context_answer_and_codes(tmp_path):
    output = tmp_path / "material.docx"
    task = SimpleNamespace(
        course_name="人工智能与创意设计",
        class_name="数媒艺术 2501",
        teacher_name="张老师",
    )
    outline = SimpleNamespace(
        session_no=1,
        topic="AIGC 与创意设计导入",
        date_text="2026-09-07",
        weekday="周一",
        periods="1-4",
    )
    material = SimpleNamespace(
        material_type="assignment",
        title="AIGC 实践作业",
        difficulty="medium",
        estimated_minutes=40,
        content="完成一份案例拆解。",
        reference_answer="成果应包含背景、方法和结论。",
        grading_criteria="任务完成度 40 分。",
        course_goal_codes="M1",
        ability_codes="1-3-4",
    )

    export_session_material_docx(output, task, outline, material)

    document = Document(output)
    text_parts = [paragraph.text for paragraph in document.paragraphs]
    text_parts.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    text = "\n".join(text_parts)
    assert "人工智能与创意设计" in text
    assert "第 1 次课" in text
    assert "参考答案" in text
    assert "评分标准" in text
    assert "M1" in text
    assert "1-3-4" in text
