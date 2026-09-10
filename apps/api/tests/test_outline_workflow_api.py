from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
import pytest
from sqlmodel import Session, select

from app.db import engine
from app.main import app
from app.models import AiModelConfig, CourseProject, CourseReviewNotice, TaskMaterialAsset
from app.routes import tasks as task_routes
from app.services.ai_outline_sections import LearningUnit, OutlineSections
from app.services.outline_generator import GeneratedOutlineRow
from app.services.secret_store import encrypt_secret


SCHOOL_OUTLINE_TEMPLATE = Path(__file__).resolve().parents[3] / "templates" / "课程实施大纲模板.docx"
SECTION_CALLS: list[str] = []


@pytest.fixture(autouse=True)
def mock_ai_outline_generation(monkeypatch):
    SECTION_CALLS.clear()

    def fake_sections(_config, evidence):
        SECTION_CALLS.append(evidence.course_name)
        return OutlineSections(
            course_summary=f"《{evidence.course_name}》带学生走完一轮创意项目。",
            teaching_strategy="线上线下混合式",
            prerequisites="设计基础",
            learning_outcomes=["能完成一个 AI 辅助的创意项目。"],
            learning_units=[
                LearningUnit(
                    name=project["name"],
                    teaching_content=project["teaching_content"],
                    teaching_methods=project["suggested_methods"],
                    key_points="重点：流程把控\n难点：提示词撰写",
                    hours=int(project["reference_hours"]),
                )
                for project in evidence.projects
            ],
            study_advice="跟着做、想着做、变着做。",
            academic_integrity="不得抄袭他人作品。",
            attendance="履行请假手续。",
            classroom_discipline="上课起立问好。",
            assignment_requirements="作业不允许迟交。",
        )

    monkeypatch.setattr("app.routes.tasks.generate_outline_sections", fake_sections, raising=False)

    def fake_generate(_config, evidence):
        project = evidence.projects[0]
        generated = []
        goal_codes = list(evidence.goals)
        for index, session in enumerate(evidence.sessions):
            goal_code = goal_codes[index % len(goal_codes)]
            goal = evidence.goals[goal_code]
            generated.append(GeneratedOutlineRow(
                session_no=session["session_no"],
                date_text=session["date_text"],
                week_no=session["week_no"],
                weekday=session["weekday"],
                periods=session["periods"],
                topic=f"第 {session['session_no']} 次课：{goal['description']}",
                teaching_content=goal["description"],
                ideological_point="原创意识与职业规范",
                teaching_methods=project["suggested_methods"],
                pre_task="预习本次课相关案例，记录主题来源和问题。",
                in_class_task="完成本次课项目任务，保留设计过程记录。",
                post_task="完善课堂成果并提交阶段材料。",
                course_goal_codes=goal_code,
                ability_codes=" ".join(goal["abilities"]),
            ))
        return generated

    monkeypatch.setattr("app.routes.tasks.generate_ai_outline", fake_generate, raising=False)
    monkeypatch.setattr(
        "app.routes.tasks.generate_outline_revision",
        lambda *_args, **_kwargs: "优化后的教学内容",
        raising=False,
    )


def auth_headers(client: TestClient) -> dict[str, str]:
    response = client.post("/auth/login", json={"employee_no": "admin", "password": "Admin@2026!"})
    assert response.status_code == 200
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def make_course_standard_docx(path: Path, with_assessment: bool = False) -> None:
    doc = Document()
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "编号"
    table.cell(0, 1).text = "课程教学目标"
    table.cell(0, 2).text = "对应的知识能力素养集"
    table.cell(1, 0).text = "M1"
    table.cell(1, 1).text = "理解AIGC基础与创意设计流程"
    table.cell(1, 2).text = "1-3-4"
    table.cell(2, 0).text = "M2"
    table.cell(2, 1).text = "完成AI辅助文创项目"
    table.cell(2, 2).text = "2-3-4 2-3-5"
    project_table = doc.add_table(rows=2, cols=7)
    headers = ["项目名称", "教学内容", "教学方法建议", "教学目标", "知识能力素养集", "知识能力素养集", "参考课时"]
    for index, header in enumerate(headers):
        project_table.cell(0, index).text = header
    values = ["项目一：AIGC 创意实践", "理解基础并完成 AI 创意项目。", "案例分析 任务驱动", "M1 M2", "1-3-4 2-3-4 2-3-5", "1-3-4 2-3-4 2-3-5", "8/4"]
    for index, value in enumerate(values):
        project_table.cell(1, index).text = value
    if with_assessment:
        doc.add_paragraph("四、课程评价")
        assessment = doc.add_table(rows=3, cols=4)
        assessment.cell(0, 0).merge(assessment.cell(0, 1)).text = "考核项目"
        assessment.cell(0, 2).text = "考核方式"
        assessment.cell(0, 3).text = "比例"
        for row_index, row in enumerate(
            [
                ("过程考核（平时成绩）", "平时表现", "按作业与课堂表现评定。", "40%"),
                ("结果考核(期末成绩)", "期末项目", "由考核小组评定作品汇报。", "60%"),
            ],
            start=1,
        ):
            for column, value in enumerate(row):
                assessment.cell(row_index, column).text = value
    doc.save(path)


def make_unknown_course_standard_docx(path: Path) -> None:
    make_course_standard_docx(path)
    doc = Document(path)
    doc.tables[0].cell(1, 2).text = "9-9-9"
    doc.save(path)


def make_talent_plan_docx(path: Path) -> None:
    doc = Document()
    table = doc.add_table(rows=3, cols=4)
    table.cell(0, 1).text = "培养规格代码"
    table.cell(0, 2).text = "TOP10"
    table.cell(0, 3).text = "其他"
    table.cell(1, 0).text = "knowledge"
    table.cell(1, 1).text = "1-3"
    table.cell(1, 2).text = "1-3-4Original keyframe indicator"
    table.cell(2, 0).text = "ability"
    table.cell(2, 1).text = "2-3"
    table.cell(2, 2).text = "2-3-4Original project indicator 2-3-5Original review indicator"
    doc.save(path)


def prepare_confirmed_sources(client: TestClient, task_id: int, standard: Path) -> dict:
    talent_plan = standard.with_name(f"{standard.stem}-talent-plan.docx")
    make_talent_plan_docx(talent_plan)

    with talent_plan.open("rb") as file:
        response = client.post(
            f"/tasks/{task_id}/talent-plan",
            files={"file": ("talent-plan.docx", file, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
    assert response.status_code == 200
    assert response.json()["indicators_count"] == 3

    with standard.open("rb") as file:
        response = client.post(
            f"/tasks/{task_id}/course-standard",
            files={"file": ("standard.docx", file, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
    assert response.status_code == 200
    assert response.json()["goals_count"] == 2
    assert response.json()["projects_count"] == 1
    with Session(engine) as session:
        projects = session.exec(select(CourseProject).where(CourseProject.task_id == task_id)).all()
    assert len(projects) == 1
    assert projects[0].reference_hours == 8

    response = client.post(f"/tasks/{task_id}/sources/confirm")
    assert response.status_code == 200
    assert response.json()["confirmed"] is True
    return response.json()


def make_schedule_xlsx(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["周次", "日期", "星期", "节次", "课程", "班级", "地点"])
    sheet.append([1, "2026-09-07", "一", "1-4", "人工智能与创意设计", "数字艺术25级1班", "智慧教室"])
    sheet.append([2, "2026-09-14", "一", "1-4", "人工智能与创意设计", "数字艺术25级1班", "智慧教室"])
    workbook.save(path)


def make_changed_schedule_xlsx(path: Path, session_count: int = 2) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["周次", "日期", "星期", "节次", "课程", "班级", "地点"])
    for index in range(1, session_count + 1):
        sheet.append([
            index,
            f"2026-09-{7 + index:02d}",
            "二",
            "5-8",
            "人工智能与创意设计",
            "数字艺术25级1班",
            "新实训室",
        ])
    workbook.save(path)


def make_outline_template(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("课程：{{课程名称}}")
    doc.add_paragraph("教师：{{任课教师}}")
    table = doc.add_table(rows=3, cols=7)
    headers = [
        "日期",
        "周次",
        "节次",
        "教学内容",
        "课程思政切入点",
        "教学方法",
        "课前、课中、课后学习要求或任务",
    ]
    for index, header in enumerate(headers):
        table.cell(0, index).text = header
    doc.save(path)


def make_lesson_template_without_heading_styles(path: Path) -> None:
    doc = Document()
    for style_name in ("Heading 1", "Heading 2"):
        style = doc.styles[style_name]
        style.element.getparent().remove(style.element)
    doc.add_paragraph("课程：{{课程名称}}")
    doc.add_paragraph("{{教案正文}}")
    doc.save(path)


def make_structured_lesson_template(path: Path, block_count: int = 2) -> None:
    doc = Document()
    for number in range(1, block_count + 1):
        info = doc.add_table(rows=8, cols=10)
        info.cell(0, 0).text = "课程名称"
        info.cell(0, 1).text = "模板课程"
        info.cell(0, 6).text = "任课教师"
        info.cell(0, 7).text = "模板教师"
        info.cell(2, 0).text = "本次课标题"
        info.cell(2, 1).text = f"模板课次 {number}"
        info.cell(2, 6).text = "授课学时"
        info.cell(2, 7).text = "4"
        info.cell(3, 0).text = "授课班级"
        info.cell(3, 1).text = "模板班级"
        info.cell(3, 4).text = "上课时间"
        info.cell(3, 5).text = "模板时间"
        info.cell(3, 7).text = "上课地点"
        info.cell(3, 8).text = "模板地点"
        info.cell(6, 0).text = "本次课教学目标"
        info.cell(6, 7).text = "课程教学目标"
        info.cell(6, 9).text = "能力指标代码"
        process = doc.add_table(rows=7, cols=6)
        for index, header in enumerate(["教学环节", "时长", "教学内容", "教师活动", "学生活动", "对应教学目标"]):
            process.cell(0, index).text = header
        for row_index, phase in enumerate(["课前", "课中", "课中", "课中", "课中", "课后"], start=1):
            process.cell(row_index, 0).text = phase
    doc.save(path)


def create_task(client: TestClient) -> int:
    client.headers.update(auth_headers(client))
    response = client.post(
        "/tasks",
        json={
            "term": "2026-2027 第一学期",
            "major": "数字媒体艺术设计",
            "class_name": "数字艺术25级1班",
            "course_name": "人工智能与创意设计",
            "teacher_name": "张明",
            "location": "智慧教室",
            "total_hours": 8,
            "hours_per_session": 4,
        },
    )
    assert response.status_code == 200
    with Session(engine) as session:
        config = session.exec(select(AiModelConfig).order_by(AiModelConfig.id)).first()
        if config is None:
            config = AiModelConfig(
                base_url="https://model.example/v1",
                model_name="outline-model",
                encrypted_api_key=encrypt_secret("sk-private"),
            )
        config.enabled = True
        config.connection_status = "connected"
        session.add(config)
        session.commit()
    return response.json()["id"]


def prepare_outline(client: TestClient, tmp_path: Path, with_assessment: bool = False) -> tuple[int, dict]:
    standard = tmp_path / "standard.docx"
    schedule = tmp_path / "schedule.xlsx"
    make_course_standard_docx(standard, with_assessment=with_assessment)
    make_schedule_xlsx(schedule)
    task_id = create_task(client)
    prepare_confirmed_sources(client, task_id, standard)
    with schedule.open("rb") as file:
        response = client.post(
            f"/tasks/{task_id}/schedule",
            files={"file": ("schedule.xlsx", file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert response.status_code == 200
    response = client.post(f"/tasks/{task_id}/outline/generate")
    assert response.status_code == 200
    return task_id, response.json()[0]


def test_outline_generation_requires_source_confirmation():
    with TestClient(app) as client:
        task_id = create_task(client)

        response = client.post(f"/tasks/{task_id}/outline/generate")

    assert response.status_code == 400
    assert "确认课程依据" in response.json()["detail"]


def test_source_review_confirmation_unlocks_outline_generation(tmp_path):
    standard = tmp_path / "standard.docx"
    talent_plan = tmp_path / "talent-plan.docx"
    schedule = tmp_path / "schedule.xlsx"
    make_course_standard_docx(standard)
    make_talent_plan_docx(talent_plan)
    make_schedule_xlsx(schedule)

    with TestClient(app) as client:
        task_id = create_task(client)

        with talent_plan.open("rb") as file:
            response = client.post(
                f"/tasks/{task_id}/talent-plan",
                files={"file": ("talent-plan.docx", file, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
        assert response.status_code == 200
        assert response.json()["indicators_count"] == 3

        with standard.open("rb") as file:
            response = client.post(
                f"/tasks/{task_id}/course-standard",
                files={"file": ("standard.docx", file, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
        assert response.status_code == 200

        response = client.get(f"/tasks/{task_id}/sources/review")
        assert response.status_code == 200
        review = response.json()
        assert review["indicators_count"] == 3
        assert review["unknown_codes"] == []
        assert review["can_confirm"] is True
        assert review["confirmed"] is False
        assert review["goals"][0]["indicators"][0]["description"] == "Original keyframe indicator"

        response = client.post(f"/tasks/{task_id}/sources/confirm")
        assert response.status_code == 200
        assert response.json()["confirmed"] is True

        with schedule.open("rb") as file:
            response = client.post(
                f"/tasks/{task_id}/schedule",
                files={"file": ("schedule.xlsx", file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        assert response.status_code == 200

        response = client.post(f"/tasks/{task_id}/outline/generate")
        second = client.post(f"/tasks/{task_id}/outline/generate")

    assert response.status_code == 200
    assert len(response.json()) == 2
    assert second.status_code == 409


def _confirmed_task_with_schedule(client: TestClient, tmp_path) -> int:
    standard = tmp_path / "standard.docx"
    talent_plan = tmp_path / "talent-plan.docx"
    schedule = tmp_path / "schedule.xlsx"
    make_course_standard_docx(standard)
    make_talent_plan_docx(talent_plan)
    make_schedule_xlsx(schedule)
    task_id = create_task(client)
    with talent_plan.open("rb") as file:
        client.post(f"/tasks/{task_id}/talent-plan", files={"file": ("talent-plan.docx", file, "application/octet-stream")})
    with standard.open("rb") as file:
        client.post(f"/tasks/{task_id}/course-standard", files={"file": ("standard.docx", file, "application/octet-stream")})
    assert client.post(f"/tasks/{task_id}/sources/confirm").status_code == 200
    with schedule.open("rb") as file:
        client.post(f"/tasks/{task_id}/schedule", files={"file": ("schedule.xlsx", file, "application/octet-stream")})
    return task_id


def test_an_outline_written_while_the_model_was_thinking_is_not_written_twice(tmp_path, monkeypatch):
    """Two clicks, one outline.

    The first request found no rows, spent minutes at the model, and outlived
    the proxy timeout; the teacher clicked again. Both wrote a full set, so the
    outline had two 第 1 次课 and the lesson run made sixteen plans for eight
    sessions. Whatever lands first is the outline; the other request must say
    so instead of adding to it.
    """
    from app.db import engine
    from app.models import OutlineRow

    original = task_routes.generate_ai_outline

    def generate_and_get_overtaken(config, evidence):
        rows = original(config, evidence)
        with Session(engine) as other:
            for row in rows:
                other.add(OutlineRow(task_id=task_id, **row.__dict__))
            other.commit()
        return rows

    with TestClient(app) as client:
        task_id = _confirmed_task_with_schedule(client, tmp_path)
        monkeypatch.setattr(task_routes, "generate_ai_outline", generate_and_get_overtaken)
        response = client.post(f"/tasks/{task_id}/outline/generate")
        listed = client.get(f"/tasks/{task_id}/outline")

    assert response.status_code == 409
    assert "已由另一次请求写入" in response.json()["detail"]
    assert [row["session_no"] for row in listed.json()] == [1, 2]


def test_a_second_generation_request_is_refused_while_the_first_is_in_flight():
    """The database check cannot see a request still at the model; the guard can."""
    from fastapi import HTTPException

    with task_routes._outline_generation_guard(42):
        with pytest.raises(HTTPException) as refused:
            with task_routes._outline_generation_guard(42):
                pass
        assert refused.value.status_code == 409
        assert "正在生成中" in refused.value.detail
        # Another course is not held up by this one.
        with task_routes._outline_generation_guard(43):
            pass
    # Released on exit, so the teacher can generate once the first run is done.
    with task_routes._outline_generation_guard(42):
        pass


def test_outline_generation_requires_enabled_model(tmp_path):
    standard = tmp_path / "standard.docx"
    schedule = tmp_path / "schedule.xlsx"
    make_course_standard_docx(standard)
    make_schedule_xlsx(schedule)

    with TestClient(app) as client:
        task_id = create_task(client)
        prepare_confirmed_sources(client, task_id, standard)
        with schedule.open("rb") as file:
            client.post(
                f"/tasks/{task_id}/schedule",
                files={"file": ("schedule.xlsx", file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        with Session(engine) as session:
            for config in session.exec(select(AiModelConfig)).all():
                session.delete(config)
            session.commit()

        response = client.post(f"/tasks/{task_id}/outline/generate")

    assert response.status_code == 409


def test_outline_update_rejects_ability_outside_selected_goals(tmp_path):
    with TestClient(app) as client:
        task_id, row = prepare_outline(client, tmp_path)
        payload = {key: value for key, value in row.items() if key not in {"id", "session_no", "post_class_recorded", "has_previous_adjustment"}}
        payload["course_goal_codes"] = "M1"
        payload["ability_codes"] = "2-3-4"

        response = client.put(f"/tasks/{task_id}/outline/{row['id']}", json=payload)

    assert response.status_code == 422
    assert "2-3-4" in response.json()["detail"]


def test_outline_revision_candidate_requires_acceptance_and_detects_stale_row(tmp_path):
    with TestClient(app) as client:
        task_id, row = prepare_outline(client, tmp_path)
        created = client.post(
            f"/tasks/{task_id}/outline/{row['id']}/revision-candidates",
            json={"field_name": "teaching_content", "instruction": "增加实践任务"},
        )
        assert created.status_code == 200
        candidate = created.json()
        assert candidate["proposed_content"] == "优化后的教学内容"

        payload = {key: value for key, value in row.items() if key not in {"id", "session_no", "post_class_recorded", "has_previous_adjustment", "updated_at"}}
        payload["teaching_content"] = "教师刚刚修改的内容"
        updated = client.put(f"/tasks/{task_id}/outline/{row['id']}", json=payload)
        assert updated.status_code == 200

        accepted = client.post(
            f"/tasks/{task_id}/outline-revision-candidates/{candidate['id']}/accept"
        )

    assert accepted.status_code == 409


def test_unknown_ability_code_blocks_source_confirmation(tmp_path):
    standard = tmp_path / "unknown-standard.docx"
    talent_plan = tmp_path / "talent-plan.docx"
    make_unknown_course_standard_docx(standard)
    make_talent_plan_docx(talent_plan)

    with TestClient(app) as client:
        task_id = create_task(client)
        with talent_plan.open("rb") as file:
            client.post(
                f"/tasks/{task_id}/talent-plan",
                files={"file": ("talent-plan.docx", file, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
        with standard.open("rb") as file:
            client.post(
                f"/tasks/{task_id}/course-standard",
                files={"file": ("standard.docx", file, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )

        review = client.get(f"/tasks/{task_id}/sources/review").json()
        response = client.post(f"/tasks/{task_id}/sources/confirm")

    assert review["unknown_codes"] == ["9-9-9"]
    assert review["can_confirm"] is False
    assert response.status_code == 400


def test_reuploading_course_standard_invalidates_confirmation(tmp_path):
    standard = tmp_path / "standard.docx"
    make_course_standard_docx(standard)

    with TestClient(app) as client:
        task_id = create_task(client)
        prepare_confirmed_sources(client, task_id, standard)

        with standard.open("rb") as file:
            response = client.post(
                f"/tasks/{task_id}/course-standard",
                files={"file": ("standard.docx", file, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
        assert response.status_code == 200

        review = client.get(f"/tasks/{task_id}/sources/review").json()
        generation = client.post(f"/tasks/{task_id}/outline/generate")

    assert review["can_confirm"] is True
    assert review["confirmed"] is False
    assert generation.status_code == 400
    assert "确认课程依据" in generation.json()["detail"]


def test_empty_source_uploads_do_not_replace_confirmed_data(tmp_path):
    standard = tmp_path / "standard.docx"
    empty_standard = tmp_path / "empty-standard.docx"
    empty_talent_plan = tmp_path / "empty-talent-plan.docx"
    make_course_standard_docx(standard)
    Document().save(empty_standard)
    Document().save(empty_talent_plan)

    with TestClient(app) as client:
        task_id = create_task(client)
        prepare_confirmed_sources(client, task_id, standard)

        with empty_standard.open("rb") as file:
            standard_response = client.post(
                f"/tasks/{task_id}/course-standard",
                files={"file": ("empty-standard.docx", file, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
        with empty_talent_plan.open("rb") as file:
            talent_response = client.post(
                f"/tasks/{task_id}/talent-plan",
                files={"file": ("empty-talent-plan.docx", file, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
        review = client.get(f"/tasks/{task_id}/sources/review").json()

    assert standard_response.status_code == 400
    assert talent_response.status_code == 400
    assert review["indicators_count"] == 3
    assert len(review["goals"]) == 2
    assert review["confirmed"] is True


def test_successful_source_upload_records_current_file_metadata(tmp_path, monkeypatch):
    monkeypatch.setattr(task_routes, "TASK_FILE_DIR", tmp_path / "task-files")
    standard = tmp_path / "standard.docx"
    make_course_standard_docx(standard)

    with TestClient(app) as client:
        task_id = create_task(client)
        with standard.open("rb") as file:
            response = client.post(
                f"/tasks/{task_id}/course-standard",
                files={"file": ("学院课程标准.docx", file, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )

    assert response.status_code == 200
    with Session(engine) as session:
        asset = session.exec(
            select(TaskMaterialAsset).where(
                TaskMaterialAsset.task_id == task_id,
                TaskMaterialAsset.kind == "course_standard",
            )
        ).one()
    assert asset.original_filename == "学院课程标准.docx"
    assert Path(asset.storage_path).exists()
    assert asset.size_bytes > 0


def test_source_upload_rejects_wrong_extension_without_replacing_current_asset(tmp_path, monkeypatch):
    monkeypatch.setattr(task_routes, "TASK_FILE_DIR", tmp_path / "task-files")
    standard = tmp_path / "standard.docx"
    make_course_standard_docx(standard)

    with TestClient(app) as client:
        task_id = create_task(client)
        with standard.open("rb") as file:
            first = client.post(
                f"/tasks/{task_id}/course-standard",
                files={"file": ("standard.docx", file, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
        with standard.open("rb") as file:
            rejected = client.post(
                f"/tasks/{task_id}/course-standard",
                files={"file": ("standard.pdf", file, "application/pdf")},
            )

    assert first.status_code == 200
    assert rejected.status_code == 400
    with Session(engine) as session:
        assets = session.exec(
            select(TaskMaterialAsset).where(
                TaskMaterialAsset.task_id == task_id,
                TaskMaterialAsset.kind == "course_standard",
            )
        ).all()
    assert len(assets) == 1
    assert assets[0].original_filename == "standard.docx"


def test_template_upload_records_active_metadata_and_rejects_wrong_extension(tmp_path, monkeypatch):
    monkeypatch.setattr(task_routes, "TASK_FILE_DIR", tmp_path / "task-files")
    template = tmp_path / "outline.docx"
    make_outline_template(template)

    with TestClient(app) as client:
        task_id = create_task(client)
        with template.open("rb") as file:
            accepted = client.post(
                f"/tasks/{task_id}/templates/outline",
                files={"file": ("学校实施大纲.docx", file, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
        with template.open("rb") as file:
            rejected = client.post(
                f"/tasks/{task_id}/templates/outline",
                files={"file": ("学校实施大纲.pdf", file, "application/pdf")},
            )

    assert accepted.status_code == 200
    assert rejected.status_code == 400
    with Session(engine) as session:
        asset = session.exec(
            select(TaskMaterialAsset).where(
                TaskMaterialAsset.task_id == task_id,
                TaskMaterialAsset.kind == "outline_template",
            )
        ).one()
    assert asset.original_filename == "学校实施大纲.docx"
    assert Path(asset.storage_path).exists()


def test_schedule_candidate_waits_for_confirmation_before_replacing_active_schedule(tmp_path, monkeypatch):
    monkeypatch.setattr(task_routes, "TASK_FILE_DIR", tmp_path / "task-files")
    original = tmp_path / "schedule.xlsx"
    changed = tmp_path / "changed.xlsx"
    make_schedule_xlsx(original)
    make_changed_schedule_xlsx(changed)

    with TestClient(app) as client:
        task_id = create_task(client)
        with original.open("rb") as file:
            assert client.post(f"/tasks/{task_id}/schedule", files={"file": ("schedule.xlsx", file)}).status_code == 200
        with changed.open("rb") as file:
            candidate = client.post(
                f"/tasks/{task_id}/schedule-candidates",
                files={"file": ("changed.xlsx", file)},
            )
        assert candidate.status_code == 200
        confirmed = client.post(
            f"/tasks/{task_id}/schedule-candidates/{candidate.json()['id']}/confirm"
        )

    assert candidate.status_code == 200
    assert candidate.json()["status"] == "pending"
    assert candidate.json()["changed_count"] == 2
    assert confirmed.status_code == 200
    with Session(engine) as session:
        records = session.exec(
            select(task_routes.ScheduleSessionRecord)
            .where(task_routes.ScheduleSessionRecord.task_id == task_id)
            .order_by(task_routes.ScheduleSessionRecord.session_no)
        ).all()
    assert records[0].date_text == "2026-09-08"
    assert records[0].periods == "5-8"


def test_schedule_candidate_with_different_count_cannot_replace_existing_outline(tmp_path, monkeypatch):
    monkeypatch.setattr(task_routes, "TASK_FILE_DIR", tmp_path / "task-files")
    changed = tmp_path / "three-sessions.xlsx"
    make_changed_schedule_xlsx(changed, session_count=3)

    with TestClient(app) as client:
        task_id, _ = prepare_outline(client, tmp_path)
        with changed.open("rb") as file:
            candidate = client.post(
                f"/tasks/{task_id}/schedule-candidates",
                files={"file": ("three-sessions.xlsx", file)},
            )
        assert candidate.status_code == 200
        response = client.post(
            f"/tasks/{task_id}/schedule-candidates/{candidate.json()['id']}/confirm"
        )

    assert response.status_code == 409
    assert "课次数量" in response.json()["detail"]


def test_course_readiness_api_returns_backend_owned_blockers():
    with TestClient(app) as client:
        task_id = create_task(client)
        response = client.get(f"/tasks/{task_id}/readiness")

    assert response.status_code == 200
    body = response.json()
    assert set(body["materials"]) == {
        "talent_plan", "course_standard", "schedule", "outline_template", "lesson_template"
    }
    assert body["can_generate_outline"] is False
    assert body["blocking_reasons"][0]["code"] == "talent_plan_missing"


def test_readiness_answers_while_a_timetable_waits_for_confirmation(tmp_path):
    """课表停在「待确认教学班」时 /readiness 仍要答得出来。

    `test_course_readiness.py` 已经断言这一步产出 `awaiting_confirmation`，但它直接
    调服务层，**不过 FastAPI 的 response_model 校验**。响应 schema 漏掉这个字面量时
    服务层测试全绿、线上却 500 —— 2026-08-11 上传课表就是这么废掉的。
    所以这条必须走 HTTP。
    """
    schedule = tmp_path / "schedule.xlsx"
    make_schedule_xlsx(schedule)

    with TestClient(app) as client:
        task_id = create_task(client)
        with schedule.open("rb") as file:
            candidate = client.post(
                f"/tasks/{task_id}/schedule-candidates",
                files={"file": ("schedule.xlsx", file)},
            )
        assert candidate.status_code == 200
        response = client.get(f"/tasks/{task_id}/readiness")

    assert response.status_code == 200
    assert response.json()["materials"]["schedule"]["status"] == "awaiting_confirmation"


def test_course_review_notice_can_be_resolved():
    with TestClient(app) as client:
        task_id = create_task(client)
        with Session(engine) as session:
            notice = CourseReviewNotice(task_id=task_id, artifact_type="outline", reason="课程标准已更新")
            session.add(notice)
            session.commit()
            session.refresh(notice)
            notice_id = notice.id
        readiness = client.get(f"/tasks/{task_id}/readiness")
        response = client.post(f"/tasks/{task_id}/review-notices/{notice_id}/resolve")
        after = client.get(f"/tasks/{task_id}/readiness")

    assert readiness.json()["review_notices"][0]["id"] == notice_id
    assert response.status_code == 200
    assert response.json()["status"] == "resolved"
    assert after.json()["review_notices"] == []


def test_replacing_authoritative_source_marks_existing_outline_for_review(tmp_path):
    standard = tmp_path / "standard.docx"
    make_course_standard_docx(standard)
    with TestClient(app) as client:
        task_id, _ = prepare_outline(client, tmp_path)
        prepare_confirmed_sources(client, task_id, standard)
        with standard.open("rb") as file:
            response = client.post(
                f"/tasks/{task_id}/course-standard",
                files={"file": ("replacement.docx", file)},
            )
        readiness = client.get(f"/tasks/{task_id}/readiness").json()

    assert response.status_code == 200
    assert any(item["artifact_type"] == "outline" for item in readiness["review_notices"])
    assert readiness["source_review"]["confirmed"] is False


def test_outline_generation_and_docx_export_workflow(tmp_path):
    standard = tmp_path / "standard.docx"
    schedule = tmp_path / "schedule.xlsx"
    template = tmp_path / "template.docx"
    make_course_standard_docx(standard)
    make_schedule_xlsx(schedule)
    make_outline_template(template)

    with TestClient(app) as client:
        task_id = create_task(client)
        prepare_confirmed_sources(client, task_id, standard)

        with schedule.open("rb") as file:
            response = client.post(
                f"/tasks/{task_id}/schedule",
                files={"file": ("schedule.xlsx", file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        assert response.status_code == 200
        assert response.json()["sessions_count"] == 2

        response = client.post(f"/tasks/{task_id}/outline/generate")
        assert response.status_code == 200
        rows = response.json()
        assert len(rows) == 2
        assert rows[0]["course_goal_codes"] == "M1"

        with template.open("rb") as file:
            response = client.post(
                f"/tasks/{task_id}/outline/export",
                files={"file": ("template.docx", file, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
        assert response.status_code == 200
        assert response.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

    output = tmp_path / "exported.docx"
    output.write_bytes(response.content)
    rendered = Document(output)
    text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
    assert "课程：人工智能与创意设计" in text
    assert "教师：张明" in text
    schedule_table = rendered.tables[0]
    assert schedule_table.cell(1, 0).text == "2026-09-07"
    assert "理解AIGC基础与创意设计流程" in schedule_table.cell(1, 3).text
    assert schedule_table.cell(2, 0).text == "2026-09-14"


def test_outline_export_uses_stored_template_when_no_file_is_uploaded(tmp_path, monkeypatch):
    monkeypatch.setattr(task_routes, "TASK_FILE_DIR", tmp_path / "task-files")
    standard = tmp_path / "standard.docx"
    schedule = tmp_path / "schedule.xlsx"
    template = tmp_path / "template.docx"
    make_course_standard_docx(standard)
    make_schedule_xlsx(schedule)
    make_outline_template(template)

    with TestClient(app) as client:
        task_id = create_task(client)

        with template.open("rb") as file:
            response = client.post(
                f"/tasks/{task_id}/templates/outline",
                files={"file": ("template.docx", file, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
        assert response.status_code == 200
        assert response.json()["kind"] == "outline"

        prepare_confirmed_sources(client, task_id, standard)
        with schedule.open("rb") as file:
            client.post(
                f"/tasks/{task_id}/schedule",
                files={"file": ("schedule.xlsx", file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        client.post(f"/tasks/{task_id}/outline/generate")

        response = client.post(f"/tasks/{task_id}/outline/export")

    assert response.status_code == 200
    output = tmp_path / "stored-template-outline.docx"
    output.write_bytes(response.content)
    rendered = Document(output)
    assert rendered.tables[0].cell(1, 0).text == "2026-09-07"


def _export_with_school_template(client: TestClient, task_id: int):
    with SCHOOL_OUTLINE_TEMPLATE.open("rb") as file:
        return client.post(
            f"/tasks/{task_id}/outline/export",
            files={
                "file": (
                    "课程实施大纲模板.docx",
                    file,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )


def test_the_school_template_comes_back_as_a_finished_outline(tmp_path):
    """Not just its schedule: 课程介绍, 学习内容, 考核方式 and the rest are filled too."""
    with TestClient(app) as client:
        task_id, _ = prepare_outline(client, tmp_path, with_assessment=True)

        response = _export_with_school_template(client, task_id)

    assert response.status_code == 200
    output = tmp_path / "school-outline.docx"
    output.write_bytes(response.content)
    rendered = Document(output)
    text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
    assert "课程简介：《人工智能与创意设计》带学生走完一轮创意项目。" in text
    assert "跟着做、想着做、变着做。" in text
    # The standard uploaded here states no resources, so the export says so
    # rather than borrowing an entry from somewhere else.
    assert "课程标准中未提供，请教师补充" in text
    assert "AI 生成" not in text

    content = next(table for table in rendered.tables if "教学单元" in table.rows[0].cells[0].text)
    assert content.cell(1, 0).text == "项目一：AIGC 创意实践"
    assert content.cell(1, 4).text == "8"
    # 六、考核方式 is the standard's own table, carried across rather than written.
    assessment = next(table for table in rendered.tables if "考核项目" in table.rows[0].cells[0].text)
    assert [row.cells[1].text for row in assessment.rows[1:]] == ["平时表现", "期末项目"]
    assert [row.cells[3].text for row in assessment.rows[1:]] == ["40%", "60%"]
    schedule = next(table for table in rendered.tables if "日期" in table.rows[0].cells[0].text)
    assert schedule.cell(1, 0).text == "2026-09-07"


def test_a_standard_that_states_no_assessment_leaves_the_split_to_the_teacher(tmp_path):
    """These percentages decide grades; the export must not invent them."""
    with TestClient(app) as client:
        task_id, _ = prepare_outline(client, tmp_path)

        response = _export_with_school_template(client, task_id)

    output = tmp_path / "no-assessment.docx"
    output.write_bytes(response.content)
    rendered = Document(output)
    assessment = next(table for table in rendered.tables if "考核项目" in table.rows[0].cells[0].text)
    assert [row.cells[0].text for row in assessment.rows[1:]] == ["课程标准中未提供，请教师补充"]


def test_the_body_is_written_once_and_then_kept(tmp_path):
    """Re-exporting to check a layout must not hand back different prose."""
    with TestClient(app) as client:
        task_id, _ = prepare_outline(client, tmp_path)

        first = _export_with_school_template(client, task_id)
        second = _export_with_school_template(client, task_id)

        assert [first.status_code, second.status_code] == [200, 200]
        assert SECTION_CALLS == ["人工智能与创意设计"]

        response = client.post(f"/tasks/{task_id}/outline/sections/regenerate")

    assert response.status_code == 200
    assert response.json()["teaching_strategy"] == "线上线下混合式"
    assert len(SECTION_CALLS) == 2


def test_a_template_that_marks_nothing_costs_no_model_call(tmp_path):
    """A teacher's own outline asks for nothing beyond 学习进程."""
    template = tmp_path / "plain-template.docx"
    make_outline_template(template)

    with TestClient(app) as client:
        task_id, _ = prepare_outline(client, tmp_path)

        with template.open("rb") as file:
            response = client.post(
                f"/tasks/{task_id}/outline/export",
                files={"file": ("template.docx", file, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )

    assert response.status_code == 200
    assert SECTION_CALLS == []


def test_lesson_export_works_when_template_lacks_heading_styles(tmp_path):
    standard = tmp_path / "standard.docx"
    schedule = tmp_path / "schedule.xlsx"
    template = tmp_path / "lesson-template.docx"
    make_course_standard_docx(standard)
    make_schedule_xlsx(schedule)
    make_lesson_template_without_heading_styles(template)

    with TestClient(app) as client:
        task_id = create_task(client)
        prepare_confirmed_sources(client, task_id, standard)
        with schedule.open("rb") as file:
            client.post(
                f"/tasks/{task_id}/schedule",
                files={"file": ("schedule.xlsx", file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        client.post(f"/tasks/{task_id}/outline/generate")
        client.post(f"/tasks/{task_id}/lessons/generate")

        with template.open("rb") as file:
            response = client.post(
                f"/tasks/{task_id}/lessons/export",
                files={"file": ("lesson-template.docx", file, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )

    assert response.status_code == 200
    output = tmp_path / "lesson-exported.docx"
    output.write_bytes(response.content)
    rendered = Document(output)
    text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
    assert "理解AIGC基础与创意设计流程" in text


def test_lesson_generation_and_docx_export_workflow(tmp_path):
    standard = tmp_path / "standard.docx"
    schedule = tmp_path / "schedule.xlsx"
    template = tmp_path / "lesson-template.docx"
    make_course_standard_docx(standard)
    make_schedule_xlsx(schedule)

    doc = Document()
    doc.add_paragraph("课程：{{课程名称}}")
    doc.add_paragraph("教案正文：")
    doc.add_paragraph("{{教案正文}}")
    doc.save(template)

    with TestClient(app) as client:
        task_id = create_task(client)
        prepare_confirmed_sources(client, task_id, standard)

        with schedule.open("rb") as file:
            response = client.post(
                f"/tasks/{task_id}/schedule",
                files={"file": ("schedule.xlsx", file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        assert response.status_code == 200

        response = client.post(f"/tasks/{task_id}/outline/generate")
        assert response.status_code == 200

        response = client.post(f"/tasks/{task_id}/lessons/generate")
        assert response.status_code == 200
        lessons = response.json()
        assert len(lessons) == 2
        assert lessons[0]["title"] == "第 1 次课：理解AIGC基础与创意设计流程"
        assert "案例分析" in lessons[0]["teaching_process"]

        response = client.get(f"/tasks/{task_id}/lessons")
        assert response.status_code == 200
        assert response.json()[0]["course_goal_codes"] == "M1"

        with template.open("rb") as file:
            response = client.post(
                f"/tasks/{task_id}/lessons/export",
                files={"file": ("lesson-template.docx", file, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
        assert response.status_code == 200
        assert response.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

    output = tmp_path / "lessons.docx"
    output.write_bytes(response.content)
    rendered = Document(output)
    text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
    assert "课程：人工智能与创意设计" in text
    assert "第 1 次课：理解AIGC基础与创意设计流程" in text
    assert "课前：预习本次课相关案例" in text


def test_lesson_export_uses_stored_template_when_no_file_is_uploaded(tmp_path, monkeypatch):
    monkeypatch.setattr(task_routes, "TASK_FILE_DIR", tmp_path / "task-files")
    standard = tmp_path / "standard.docx"
    schedule = tmp_path / "schedule.xlsx"
    template = tmp_path / "lesson-template.docx"
    make_course_standard_docx(standard)
    make_schedule_xlsx(schedule)

    doc = Document()
    doc.add_paragraph("课程：{{课程名称}}")
    doc.add_paragraph("{{教案正文}}")
    doc.save(template)

    with TestClient(app) as client:
        task_id = create_task(client)

        with template.open("rb") as file:
            response = client.post(
                f"/tasks/{task_id}/templates/lesson",
                files={"file": ("lesson-template.docx", file, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
        assert response.status_code == 200

        prepare_confirmed_sources(client, task_id, standard)
        with schedule.open("rb") as file:
            client.post(
                f"/tasks/{task_id}/schedule",
                files={"file": ("schedule.xlsx", file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        client.post(f"/tasks/{task_id}/outline/generate")
        client.post(f"/tasks/{task_id}/lessons/generate")

        response = client.post(f"/tasks/{task_id}/lessons/export")

    assert response.status_code == 200
    output = tmp_path / "stored-template-lessons.docx"
    output.write_bytes(response.content)
    rendered = Document(output)
    text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
    assert "第 1 次课：理解AIGC基础与创意设计流程" in text


def test_lesson_export_fills_structured_template_with_schedule_context(tmp_path):
    standard = tmp_path / "standard.docx"
    schedule = tmp_path / "schedule.xlsx"
    template = tmp_path / "structured-lesson-template.docx"
    make_course_standard_docx(standard)
    make_schedule_xlsx(schedule)
    make_structured_lesson_template(template)

    with TestClient(app) as client:
        task_id = create_task(client)
        prepare_confirmed_sources(client, task_id, standard)
        with schedule.open("rb") as file:
            client.post(
                f"/tasks/{task_id}/schedule",
                files={"file": ("schedule.xlsx", file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        client.post(f"/tasks/{task_id}/outline/generate")
        client.post(f"/tasks/{task_id}/lessons/generate")
        with template.open("rb") as file:
            response = client.post(
                f"/tasks/{task_id}/lessons/export",
                files={"file": ("lesson-template.docx", file, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )

    assert response.status_code == 200
    output = tmp_path / "structured-lessons.docx"
    output.write_bytes(response.content)
    rendered = Document(output)
    assert rendered.tables[0].cell(2, 1).text.startswith("第 1 次课：")
    assert rendered.tables[0].cell(3, 5).text == "第 1 周 周一 第 1-4 节"
    assert rendered.tables[0].cell(3, 8).text == "智慧教室"
    assert rendered.tables[2].cell(2, 1).text.startswith("第 2 次课：")
    assert rendered.tables[2].cell(3, 5).text == "第 2 周 周一 第 1-4 节"


def test_lesson_export_rejects_unrecognized_template(tmp_path):
    standard = tmp_path / "standard.docx"
    schedule = tmp_path / "schedule.xlsx"
    template = tmp_path / "unrecognized.docx"
    make_course_standard_docx(standard)
    make_schedule_xlsx(schedule)
    doc = Document()
    doc.add_paragraph("普通教案模板")
    doc.save(template)

    with TestClient(app) as client:
        task_id = create_task(client)
        prepare_confirmed_sources(client, task_id, standard)
        with schedule.open("rb") as file:
            client.post(
                f"/tasks/{task_id}/schedule",
                files={"file": ("schedule.xlsx", file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        client.post(f"/tasks/{task_id}/outline/generate")
        client.post(f"/tasks/{task_id}/lessons/generate")
        with template.open("rb") as file:
            response = client.post(
                f"/tasks/{task_id}/lessons/export",
                files={"file": ("lesson-template.docx", file, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )

    assert response.status_code == 400
    assert "未识别" in response.json()["detail"]


def test_task_list_returns_material_and_generation_progress(tmp_path, monkeypatch):
    monkeypatch.setattr(task_routes, "TASK_FILE_DIR", tmp_path / "task-files")
    standard = tmp_path / "standard.docx"
    schedule = tmp_path / "schedule.xlsx"
    outline_template = tmp_path / "outline-template.docx"
    lesson_template = tmp_path / "lesson-template.docx"
    make_course_standard_docx(standard)
    make_schedule_xlsx(schedule)
    make_outline_template(outline_template)

    doc = Document()
    doc.add_paragraph("课程：{{课程名称}}")
    doc.add_paragraph("{{教案正文}}")
    doc.save(lesson_template)

    with TestClient(app) as client:
        task_id = create_task(client)

        with outline_template.open("rb") as file:
            client.post(
                f"/tasks/{task_id}/templates/outline",
                files={"file": ("outline-template.docx", file, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
        with lesson_template.open("rb") as file:
            client.post(
                f"/tasks/{task_id}/templates/lesson",
                files={"file": ("lesson-template.docx", file, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
        prepare_confirmed_sources(client, task_id, standard)
        with schedule.open("rb") as file:
            client.post(
                f"/tasks/{task_id}/schedule",
                files={"file": ("schedule.xlsx", file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        client.post(f"/tasks/{task_id}/outline/generate")
        client.post(f"/tasks/{task_id}/lessons/generate")

        response = client.get("/tasks")

    assert response.status_code == 200
    task = next(item for item in response.json() if item["id"] == task_id)
    assert task["course_standard_uploaded"] is True
    assert task["talent_plan_uploaded"] is True
    assert task["sources_confirmed"] is True
    assert task["schedule_uploaded"] is True
    assert task["outline_template_uploaded"] is True
    assert task["lesson_template_uploaded"] is True
    assert task["outline_rows_count"] == 2
    assert task["lesson_plans_count"] == 2
    assert task["session_count"] == 2
    assert task["completed_sessions_count"] >= 0
    assert "next_session_no" in task
    assert "next_session_date" in task
    assert "next_session_weekday" in task
    assert "next_session_periods" in task


def test_course_standard_without_project_table_is_rejected_at_upload(tmp_path):
    # Marking a goals-only standard "ready" and failing at outline generation
    # sent teachers hunting in the wrong place; reject it with the fix spelled out.
    standard = tmp_path / "goals-only-standard.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "编号"
    table.cell(0, 1).text = "课程教学目标"
    table.cell(0, 2).text = "对应的知识能力素养集"
    table.cell(1, 0).text = "M1"
    table.cell(1, 1).text = "理解基础"
    table.cell(1, 2).text = "1-3-4"
    doc.save(standard)

    with TestClient(app) as client:
        task_id = create_task(client)
        with standard.open("rb") as file:
            response = client.post(
                f"/tasks/{task_id}/course-standard",
                files={"file": ("standard.docx", file, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
        readiness = client.get(f"/tasks/{task_id}/readiness").json()

    assert response.status_code == 400
    assert "参考课时" in response.json()["detail"]
    assert readiness["materials"]["course_standard"]["status"] == "missing"


def test_outline_export_prints_the_owners_profile_in_the_teacher_block(tmp_path, monkeypatch):
    """The school template leaves 「教师姓名：」「办公地点：」… blank; the export
    writes the course owner's profile there and the class and room from the
    course record, and leaves the biography blank rather than inventing one
    when the profile is missing."""
    monkeypatch.setattr(task_routes, "TASK_FILE_DIR", tmp_path / "task-files")
    with TestClient(app) as client:
        client.headers.update(auth_headers(client))
        task_id, _ = prepare_outline(client, tmp_path)
        saved = client.put(
            "/auth/profile",
            json={"office_location": "信息楼316", "phone": "13800000000", "bio": "讲师，研究方向为数字媒体。"},
        )
        assert saved.status_code == 200

        response = _export_with_school_template(client, task_id)
        assert response.status_code == 200

    output = tmp_path / "with-profile.docx"
    output.write_bytes(response.content)
    text = [paragraph.text for paragraph in Document(output).paragraphs]
    assert "办公地点：信息楼316" in text
    assert "联系电话：13800000000" in text
    assert "教师简介：讲师，研究方向为数字媒体。" in text
    assert any(line.startswith("教师姓名：") and len(line) > len("教师姓名：") for line in text)
    assert any(line.startswith("上课班级：") and len(line) > len("上课班级：") for line in text)

