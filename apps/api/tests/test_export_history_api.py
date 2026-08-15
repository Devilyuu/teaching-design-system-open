from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
import pytest
from sqlmodel import Session

from app.db import engine
from app.main import app
from app.models import LessonPlan, OutlineRow
from app.routes import tasks as task_routes


def auth_headers(client: TestClient) -> dict[str, str]:
    response = client.post("/auth/login", json={"employee_no": "admin", "password": "Admin@2026!"})
    assert response.status_code == 200
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def create_task(client: TestClient) -> int:
    client.headers.update(auth_headers(client))
    response = client.post(
        "/tasks",
        json={
            "term": "2026-2027 第一学期",
            "major": "数字媒体艺术设计",
            "class_name": "数字艺术25级1班",
            "course_name": "人工智能与创意设计",
            "teacher_name": "张明",
            "location": "智慧教室",
            "total_hours": 8,
            "hours_per_session": 4,
        },
    )
    assert response.status_code == 200
    return response.json()["id"]


def make_outline_template(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("课程：{{课程名称}}")
    table = doc.add_table(rows=2, cols=7)
    headers = ["日期", "周次", "节次", "教学内容", "课程思政切入点", "教学方法", "课前、课中、课后学习要求或任务"]
    for index, header in enumerate(headers):
        table.cell(0, index).text = header
    for index in range(7):
        table.cell(1, index).text = "模板示例"
    doc.save(path)


def make_lesson_template(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("{{教案正文}}")
    doc.save(path)


def seed_outline_row(task_id: int) -> int:
    with Session(engine) as session:
        row = OutlineRow(
            task_id=task_id,
            session_no=1,
            date_text="2026-09-07",
            week_no=1,
            weekday="一",
            periods="1-4",
            topic="AIGC 与创意设计导入",
            teaching_content="理解 AIGC 基础与创意设计流程",
            ideological_point="技术向善",
            teaching_methods="案例分析",
            pre_task="预习",
            in_class_task="完成案例拆解",
            post_task="提交记录",
            course_goal_codes="M1",
            ability_codes="1-3-4",
        )
        session.add(row)
        session.commit()
        session.refresh(row)
        return row.id


def seed_lesson(task_id: int, outline_row_id: int) -> None:
    with Session(engine) as session:
        session.add(
            LessonPlan(
                task_id=task_id,
                outline_row_id=outline_row_id,
                session_no=1,
                title="第 1 次课",
                teaching_goals="理解 AIGC 基础",
                key_points="AIGC 流程",
                difficult_points="工具选型",
                teaching_process="导入、讲授、实训",
                homework="提交记录",
                course_goal_codes="M1",
                ability_codes="1-3-4",
            )
        )
        session.commit()


@pytest.fixture(autouse=True)
def isolate_task_files(tmp_path, monkeypatch):
    monkeypatch.setattr(task_routes, "TASK_FILE_DIR", tmp_path / "task-files")


def upload_template(client: TestClient, task_id: int, kind: str, path: Path) -> None:
    with path.open("rb") as file:
        response = client.post(f"/tasks/{task_id}/templates/{kind}", files={"file": (f"{kind}-模板.docx", file)})
    assert response.status_code == 200


def test_outline_export_is_recorded_with_the_template_it_used(tmp_path):
    template = tmp_path / "outline-template.docx"
    make_outline_template(template)

    with TestClient(app) as client:
        task_id = create_task(client)
        seed_outline_row(task_id)
        upload_template(client, task_id, "outline", template)
        exported = client.post(f"/tasks/{task_id}/outline/export")
        assert exported.status_code == 200
        history = client.get(f"/tasks/{task_id}/exports")

    assert history.status_code == 200
    records = history.json()
    assert len(records) == 1
    record = records[0]
    assert record["artifact_type"] == "outline"
    assert record["artifact_label"] == "课程实施大纲"
    assert record["filename"] == "人工智能与创意设计_课程实施大纲.docx"
    assert record["template_filename"] == "outline-模板.docx"
    assert record["size_bytes"] == len(exported.content)
    assert record["exported_by"] == "系统管理员"
    assert record["available"] is True
    assert "1 行" in record["source_summary"]


def test_exported_document_can_be_downloaded_again(tmp_path):
    template = tmp_path / "outline-template.docx"
    make_outline_template(template)

    with TestClient(app) as client:
        task_id = create_task(client)
        seed_outline_row(task_id)
        upload_template(client, task_id, "outline", template)
        exported = client.post(f"/tasks/{task_id}/outline/export")
        record_id = client.get(f"/tasks/{task_id}/exports").json()[0]["id"]
        again = client.get(f"/tasks/{task_id}/exports/{record_id}/download")

    assert again.status_code == 200
    assert again.content == exported.content
    assert again.headers["content-disposition"] == exported.headers["content-disposition"]
    assert again.headers["content-disposition"].startswith("attachment; filename*=UTF-8''")


def test_lesson_export_is_recorded_with_the_lesson_template(tmp_path):
    outline_template = tmp_path / "outline-template.docx"
    lesson_template = tmp_path / "lesson-template.docx"
    make_outline_template(outline_template)
    make_lesson_template(lesson_template)

    with TestClient(app) as client:
        task_id = create_task(client)
        row_id = seed_outline_row(task_id)
        seed_lesson(task_id, row_id)
        upload_template(client, task_id, "lesson", lesson_template)
        exported = client.post(f"/tasks/{task_id}/lessons/export")
        assert exported.status_code == 200
        records = client.get(f"/tasks/{task_id}/exports").json()

    assert len(records) == 1
    assert records[0]["artifact_type"] == "lesson"
    assert records[0]["artifact_label"] == "整门课教案"
    assert records[0]["template_filename"] == "lesson-模板.docx"
    assert "1 份" in records[0]["source_summary"]


def test_history_lists_the_newest_export_first(tmp_path):
    outline_template = tmp_path / "outline-template.docx"
    lesson_template = tmp_path / "lesson-template.docx"
    make_outline_template(outline_template)
    make_lesson_template(lesson_template)

    with TestClient(app) as client:
        task_id = create_task(client)
        row_id = seed_outline_row(task_id)
        seed_lesson(task_id, row_id)
        upload_template(client, task_id, "outline", outline_template)
        upload_template(client, task_id, "lesson", lesson_template)
        assert client.post(f"/tasks/{task_id}/outline/export").status_code == 200
        assert client.post(f"/tasks/{task_id}/lessons/export").status_code == 200
        records = client.get(f"/tasks/{task_id}/exports").json()

    assert [item["artifact_type"] for item in records] == ["lesson", "outline"]


def test_repeated_exports_are_kept_as_separate_versions(tmp_path):
    template = tmp_path / "outline-template.docx"
    make_outline_template(template)

    with TestClient(app) as client:
        task_id = create_task(client)
        seed_outline_row(task_id)
        upload_template(client, task_id, "outline", template)
        assert client.post(f"/tasks/{task_id}/outline/export").status_code == 200
        assert client.post(f"/tasks/{task_id}/outline/export").status_code == 200
        records = client.get(f"/tasks/{task_id}/exports").json()
        downloads = [
            client.get(f"/tasks/{task_id}/exports/{item['id']}/download") for item in records
        ]

    assert len(records) == 2
    assert records[0]["id"] != records[1]["id"]
    assert all(item.status_code == 200 for item in downloads)


def test_download_does_not_cross_courses(tmp_path):
    template = tmp_path / "outline-template.docx"
    make_outline_template(template)

    with TestClient(app) as client:
        task_id = create_task(client)
        other_task_id = create_task(client)
        seed_outline_row(task_id)
        upload_template(client, task_id, "outline", template)
        assert client.post(f"/tasks/{task_id}/outline/export").status_code == 200
        record_id = client.get(f"/tasks/{task_id}/exports").json()[0]["id"]
        response = client.get(f"/tasks/{other_task_id}/exports/{record_id}/download")

    assert response.status_code == 404


def test_failed_export_leaves_no_history_entry(tmp_path):
    with TestClient(app) as client:
        task_id = create_task(client)
        seed_outline_row(task_id)
        response = client.post(f"/tasks/{task_id}/outline/export")
        records = client.get(f"/tasks/{task_id}/exports").json()

    assert response.status_code == 400
    assert records == []
