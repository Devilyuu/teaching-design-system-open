from docx import Document

from app.services.course_standard_resources import (
    MISSING,
    extract_course_resources,
    read_course_resources,
    resource_lines,
)


def _standard(path, lines):
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    doc.save(path)
    return path


def test_copies_each_kind_of_resource_the_standard_states(tmp_path):
    path = _standard(
        tmp_path / "standard.docx",
        [
            "四、课程评价",
            "五、课程教学资源",
            "1．实践条件",
            "机房一间",
            "2．师资条件",
            "任课教师应具备中级及以上职称",
            "3．参考教材及教学参考书建议",
            "（1）教材",
            "（按要求选用国家规划教材）",
            "移动UI界面设计（微课版），张晓景/李晓斌，人民邮电出版社（2018）",
            "（2）参考书",
            "APP草图+流程图+交互原型设计教程，刘源，电子工业出版社（2020）",
            "PS APP UI设计从零开始学，贾浩梅，清华大学出版社（2022）",
            "4．学习资源选用",
            "（网络资源、多媒体教学课件等教学资源下载地址）",
            "https://www.zcool.com.cn/",
            "六、编制说明",
            "编写：张明",
        ],
    )

    resources = extract_course_resources(path)

    assert resources.textbooks == ["移动UI界面设计（微课版），张晓景/李晓斌，人民邮电出版社（2018）"]
    assert resources.references == [
        "APP草图+流程图+交互原型设计教程，刘源，电子工业出版社（2020）",
        "PS APP UI设计从零开始学，贾浩梅，清华大学出版社（2022）",
    ]
    assert resources.online == ["https://www.zcool.com.cn/"]
    # 实践条件、师资条件属于同一节但不是资源，不能混进来
    assert "机房一间" not in resources.textbooks + resources.references + resources.online
    assert not any("职称" in line for line in resources.references)
    # 括号里的填写说明不是资源
    assert not any(line.startswith("（按要求") for line in resources.textbooks)
    # 下一节开始就停止
    assert not any("张明" in line for line in resources.online)


def test_a_standard_without_a_resource_section_yields_nothing(tmp_path):
    path = _standard(tmp_path / "bare.docx", ["一、课程定位", "二、课程目标"])

    resources = extract_course_resources(path)

    assert resources.found_anything is False


def test_a_subsection_the_standard_omits_says_so_rather_than_borrowing(tmp_path):
    """Padding an empty subsection from a neighbour is the same defect as inventing it."""
    doc = Document()
    for line in ["五、课程教学资源", "3．参考教材及教学参考书建议", "（1）教材", "自编讲义"]:
        doc.add_paragraph(line)

    lines = resource_lines(read_course_resources(doc))

    assert lines["教材"] == ["自编讲义"]
    assert lines["教辅"] == [MISSING]
    assert lines["在线学习资源"] == [MISSING]
    # 课程标准里从来没有「高质量作业范例」这一项，永远交给教师
    assert lines["高质量作业范例"] == [MISSING]


def test_a_label_and_its_entry_on_one_line_are_read_too(tmp_path):
    """Some standards skip the （1）/（2） sub-headings and write 「教材：xxx」."""
    path = _standard(
        tmp_path / "inline.docx",
        [
            "五、课程教学资源",
            "1．实践条件",
            "机房一间",
            "3．参考教材及教学参考书建议",
            "教材：After Effects 从入门到精通，李某编著，某大学出版社",
            "参考书：《动态图形设计》倪某编著，某美术出版社，2022年12月",
            "4．学习资源选用",
            "课程资源库：",
            "http://example.invalid/course",
            "六、编制说明",
        ],
    )

    resources = extract_course_resources(path)

    assert resources.textbooks == ["After Effects 从入门到精通，李某编著，某大学出版社"]
    assert resources.references == ["《动态图形设计》倪某编著，某美术出版社，2022年12月"]
    assert resources.online == ["课程资源库：", "http://example.invalid/course"]
