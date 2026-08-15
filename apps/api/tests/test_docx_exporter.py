from types import SimpleNamespace

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.services.docx_exporter import (
    OutlineTemplateError,
    fill_docx_placeholders,
    fill_lesson_docx,
    fill_outline_docx,
)
from app.services.lesson_template_filler import LessonTemplateError


def test_fill_docx_placeholders_in_paragraphs_and_tables(tmp_path):
    template = tmp_path / "template.docx"
    output = tmp_path / "output.docx"

    doc = Document()
    doc.add_paragraph("课程：{{课程名称}}")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "任课教师"
    table.cell(0, 1).text = "{{任课教师}}"
    doc.save(template)

    fill_docx_placeholders(
        template,
        output,
        {"课程名称": "人工智能与创意设计", "任课教师": "张明"},
    )

    rendered = Document(output)
    assert rendered.paragraphs[0].text == "课程：人工智能与创意设计"
    assert rendered.tables[0].cell(0, 1).text == "张明"


def test_fill_outline_docx_detects_and_populates_schedule_table(tmp_path):
    template = tmp_path / "template.docx"
    output = tmp_path / "output.docx"

    doc = Document()
    table = doc.add_table(rows=3, cols=7)
    headers = [
        "日期",
        "周次",
        "节次",
        "教学内容",
        "课程思政切入点",
        "教学方法",
        "课前、课中、课后学习要求或任务",
    ]
    for index, header in enumerate(headers):
        table.cell(0, index).text = header
    for row_index in (1, 2):
        for column_index in range(7):
            table.cell(row_index, column_index).text = "模板示例"
    doc.save(template)

    fill_outline_docx(
        template,
        output,
        {"课程名称": "人工智能与创意设计"},
        [
            {
                "date_text": "2026-09-07",
                "week_no": 1,
                "periods": "1-4",
                "teaching_content": "AIGC基础与创意设计流程",
                "ideological_point": "技术向善",
                "teaching_methods": "案例分析、任务驱动",
                "pre_task": "了解一个文创IP案例",
                "in_class_task": "完成项目方向初选",
                "post_task": "提交案例分析",
            }
        ],
    )

    rendered = Document(output)
    schedule = rendered.tables[0]
    assert schedule.cell(1, 0).text == "2026-09-07"
    assert schedule.cell(1, 1).text == "1"
    assert schedule.cell(1, 2).text == "1-4"
    assert schedule.cell(1, 3).text == "AIGC基础与创意设计流程"
    assert schedule.cell(1, 4).text == "技术向善"
    assert schedule.cell(1, 5).text == "案例分析、任务驱动"
    assert schedule.cell(1, 6).text == (
        "课前：了解一个文创IP案例\n"
        "课中：完成项目方向初选\n"
        "课后：提交案例分析"
    )
    assert len(schedule.cell(1, 6).paragraphs) == 1
    # Surplus template rows are removed, not blanked: blanking left the
    # template's own text standing in every unmapped column.
    assert len(schedule.rows) == 2


def test_outline_schedule_columns_share_the_width_the_prose_needs(tmp_path):
    """The template sizes 课程思政 for "课堂纪律要求"; a generated row writes a sentence."""
    from docx.shared import Twips

    template = tmp_path / "template.docx"
    output = tmp_path / "output.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=7)
    headers = ["日期", "周次", "节次", "教学内容", "课程思政切入点", "教学方法", "关键教学要求或任务"]
    drawn = [816, 559, 618, 3326, 709, 739, 2228]
    for index, (header, width) in enumerate(zip(headers, drawn)):
        table.cell(0, index).text = header
        table.columns[index].width = Twips(width)
        for row_index in range(2):
            table.cell(row_index, index).width = Twips(width)
    doc.save(template)

    fill_outline_docx(template, output, {}, [
        {
            "date_text": "2026-03-02",
            "week_no": 1,
            "periods": "1-4",
            "teaching_content": "使用软件设计并发放调研问卷，完成APP市场调研",
            "ideological_point": "培养实事求是、严谨调研的科学态度",
            "teaching_methods": "讲授；讨论；练习；分组合作",
            "in_class_task": "完成并发布调研问卷，收集数据",
        },
    ])

    grid = Document(output).tables[0]._tbl.find(qn("w:tblGrid"))
    widths = [int(column.get(qn("w:w"))) for column in grid.findall(qn("w:gridCol"))]
    # 日期/周次/节次 hold fixed-format values and keep the width they were drawn
    # at; the four prose columns divide what they already occupied, 30:21:21:28.
    assert widths == [816, 559, 618, 2101, 1470, 1470, 1961]
    assert sum(widths) == sum(drawn)


def _outline_template(path, *, declared_course: str | None, in_table: bool = False):
    doc = Document()
    doc.add_paragraph("示例职业技术学院课程实施大纲")
    doc.add_paragraph("一、基本信息")
    if declared_course is not None and not in_table:
        doc.add_paragraph(f"课程名称：{declared_course}")
    if in_table:
        info = doc.add_table(rows=1, cols=2)
        info.cell(0, 0).text = f"课程名称：{declared_course}"
        info.cell(0, 1).text = "课程编码：12012092040"
    table = doc.add_table(rows=2, cols=7)
    for index, header in enumerate(
        ["日期", "周次", "节次", "教学内容", "课程思政切入点", "教学方法", "关键教学要求或任务"]
    ):
        table.cell(0, index).text = header
    doc.save(path)


def _one_row():
    return [{"date_text": "2026-03-02", "week_no": 1, "periods": "1-4", "teaching_content": "市场调研"}]


def test_an_outline_template_written_for_another_course_is_refused(tmp_path):
    """Only 学习进程 is generated; the rest of the template ships as written.

    A teacher who starts from a colleague's finished outline would otherwise get
    a complete, plausible document whose course introduction, objectives and
    assessment all belong to a different course.
    """
    template = tmp_path / "other-course.docx"
    _outline_template(template, declared_course="人工智能与创意设计")

    with pytest.raises(OutlineTemplateError) as refusal:
        fill_outline_docx(template, tmp_path / "out.docx", {"课程名称": "移动终端APP设计"}, _one_row())

    assert "人工智能与创意设计" in str(refusal.value)
    assert "移动终端APP设计" in str(refusal.value)
    assert not (tmp_path / "out.docx").exists()


def test_the_course_name_is_also_found_inside_the_information_table(tmp_path):
    """One school template puts it in a paragraph, another in a table cell."""
    template = tmp_path / "other-course-table.docx"
    _outline_template(template, declared_course="人工智能与创意设计", in_table=True)

    with pytest.raises(OutlineTemplateError, match="人工智能与创意设计"):
        fill_outline_docx(template, tmp_path / "out.docx", {"课程名称": "移动终端APP设计"}, _one_row())


def test_a_blank_template_is_the_normal_way_to_start_and_passes(tmp_path):
    """Headings and empty tables, no course named -- nothing to conflict with."""
    template = tmp_path / "blank.docx"
    _outline_template(template, declared_course="")
    output = tmp_path / "out.docx"

    fill_outline_docx(template, output, {"课程名称": "移动终端APP设计"}, _one_row())

    assert Document(output).tables[0].cell(1, 0).text == "2026-03-02"


def test_a_template_for_this_same_course_passes(tmp_path):
    template = tmp_path / "same-course.docx"
    _outline_template(template, declared_course="移动终端APP设计")
    output = tmp_path / "out.docx"

    fill_outline_docx(template, output, {"课程名称": "移动终端APP设计"}, _one_row())

    assert output.exists()


def test_a_placeholder_template_resolves_before_the_check_runs(tmp_path):
    """{{课程名称}} becomes this course, so it must not read as another one."""
    template = tmp_path / "placeholder.docx"
    _outline_template(template, declared_course="{{课程名称}}")
    output = tmp_path / "out.docx"

    fill_outline_docx(template, output, {"课程名称": "移动终端APP设计"}, _one_row())

    assert output.exists()


def _make_structured_lesson_template(path):
    doc = Document()
    info = doc.add_table(rows=8, cols=10)
    info.cell(0, 0).text = "课程名称"
    info.cell(0, 1).text = "模板课程"
    info.cell(0, 6).text = "任课教师"
    info.cell(0, 7).text = "模板教师"
    info.cell(2, 0).text = "本次课标题"
    info.cell(2, 1).text = "模板课次"
    info.cell(2, 6).text = "授课学时"
    info.cell(2, 7).text = "4"
    info.cell(3, 0).text = "授课班级"
    info.cell(3, 1).text = "模板班级"
    info.cell(3, 4).text = "上课时间"
    info.cell(3, 5).text = "模板时间"
    info.cell(3, 7).text = "上课地点"
    info.cell(3, 8).text = "模板地点"
    info.cell(6, 0).text = "本次课教学目标"
    info.cell(6, 7).text = "课程教学目标"
    info.cell(6, 9).text = "能力指标代码"
    info.cell(7, 0).text = "模板目标"
    info.cell(7, 7).text = "M1"
    info.cell(7, 9).text = "1-1-1"
    process = doc.add_table(rows=7, cols=6)
    for index, header in enumerate(["教学环节", "时长", "教学内容", "教师活动", "学生活动", "对应教学目标"]):
        process.cell(0, index).text = header
    for row_index, phase in enumerate(["课前", "课中", "课中", "课中", "课中", "课后"], start=1):
        process.cell(row_index, 0).text = phase
    doc.save(path)


def _export_lesson():
    return SimpleNamespace(
        title="第 1 次课：AIGC 导入",
        duration_minutes=160,
        teaching_goals="理解 AIGC 基础",
        teaching_process="课前：阅读案例\n导入：分析案例\n课后：完善作品",
        homework="提交作品",
        course_goal_codes="M1",
        ability_codes="1-3-4",
    )


def test_fill_lesson_docx_uses_structured_blocks_without_appending_plain_text(tmp_path):
    template = tmp_path / "lesson-template.docx"
    output = tmp_path / "lesson-output.docx"
    _make_structured_lesson_template(template)

    fill_lesson_docx(
        template,
        output,
        {"课程名称": "人工智能与创意设计", "任课教师": "张老师", "授课班级": "数媒 2501"},
        [_export_lesson()],
        contexts=[{"week_no": 1, "weekday": "周一", "periods": "1-4", "location": "A101"}],
    )

    rendered = Document(output)
    assert rendered.tables[0].cell(2, 1).text == "第 1 次课：AIGC 导入"
    assert rendered.tables[0].cell(3, 5).text == "第 1 周 周一 第 1-4 节"
    assert rendered.tables[1].cell(1, 2).text == "阅读案例"
    assert rendered.tables[1].cell(6, 2).text == "课后巩固与提交"
    assert rendered.tables[1].cell(6, 3).text == "提交作品"
    assert "教案正文" not in "\n".join(paragraph.text for paragraph in rendered.paragraphs)


def test_fill_lesson_docx_rejects_unrecognized_template_without_placeholder(tmp_path):
    template = tmp_path / "unrecognized.docx"
    output = tmp_path / "lesson-output.docx"
    doc = Document()
    doc.add_paragraph("普通教案模板")
    doc.save(template)

    with pytest.raises(LessonTemplateError, match="未识别"):
        fill_lesson_docx(template, output, {}, [_export_lesson()], contexts=[{}])


def test_generated_cells_match_the_font_of_the_table(tmp_path):
    """An empty template cell must not come back a size larger than its row."""
    from docx.shared import Pt

    template = tmp_path / "template.docx"
    output = tmp_path / "output.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=4)
    for index, header in enumerate(["日期", "周次", "节次", "教学内容"]):
        table.cell(0, index).text = header
    for index in range(4):
        cell = table.cell(1, index)
        cell.text = ""
        properties = cell.paragraphs[0]._p.get_or_add_pPr()
        run_properties = OxmlElement("w:rPr")
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), "18")  # 9pt, expressed in half-points
        run_properties.append(size)
        properties.append(run_properties)
    doc.save(template)

    fill_outline_docx(template, output, {}, [
        {"date_text": "2026-03-02", "week_no": 1, "periods": "1-4", "teaching_content": "调研问卷设计"},
    ])

    rendered = Document(output)
    sizes = {
        run.font.size
        for column in range(4)
        for run in rendered.tables[0].cell(1, column).paragraphs[0].runs
        if run.text.strip()
    }
    assert sizes == {Pt(9)}


def test_generated_text_copies_the_format_the_template_marks_in_red(tmp_path):
    """The 黑体 labels outnumber the content; the red marking is what counts.

    Word splits "授课班级" into four one-character runs, so counting runs put the
    label font ahead of the body font and the whole table came out in 黑体.
    """
    from docx.shared import Pt, RGBColor

    template = tmp_path / "lesson-template.docx"
    output = tmp_path / "lesson-output.docx"
    _make_structured_lesson_template(template)

    doc = Document(template)
    info = doc.tables[0]
    for row in info.rows:
        for cell in row.cells:
            for run in cell.paragraphs[0].runs:
                run.font.name = "黑体"
                run.font.size = Pt(10.5)

    def mark(cell, text, points):
        cell.paragraphs[0].runs[0].text = ""
        run = cell.paragraphs[0].add_run(text)
        run.font.name = "宋体"
        run.font.size = Pt(points)
        run.font.color.rgb = RGBColor.from_string("C00000")

    # Two marked sizes: the larger one has more runs, the smaller more text.
    mark(info.cell(2, 1), "模板课次的示例说明文字", 10.5)
    mark(info.cell(3, 1), "模板班级", 11)
    mark(info.cell(3, 5), "模板时间", 11)
    doc.save(template)

    fill_lesson_docx(
        template,
        output,
        {"课程名称": "人工智能与创意设计", "任课教师": "张老师", "授课班级": "数媒 2501"},
        [_export_lesson()],
        contexts=[{"week_no": 1, "weekday": "周一", "periods": "1-4", "location": "A101"}],
    )

    written = Document(output).tables[0].cell(2, 1).paragraphs[0].runs[0]
    assert written.text == "第 1 次课：AIGC 导入"
    assert written.font.name == "宋体"
    assert written.font.size == Pt(10.5)
    # The red only marked what to generate; the document itself is black.
    assert written.font.color.rgb is None
