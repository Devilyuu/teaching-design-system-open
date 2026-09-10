"""Word paints list numbers that never reach paragraph.text; the parsers must
see the same "1-1-1" a teacher sees in the printed plan."""

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.services.docx_numbering import cell_text, paragraph_numbers
from app.services.talent_plan_parser import parse_talent_plan


def add_multilevel_list(doc: Document, *, starts: tuple[int, int, int] = (1, 1, 1), override_level1: int | None = None) -> str:
    """Define a %1 / %1-%2 / %1-%2-%3 list like the school's plans and return its numId."""
    numbering = doc.part.numbering_part.element
    existing = [int(item.get(qn("w:abstractNumId"))) for item in numbering.findall(qn("w:abstractNum"))]
    abstract_id = str(max(existing, default=-1) + 1)
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), abstract_id)
    for ilvl, (start, text) in enumerate(zip(starts, ("%1", "%1-%2", "%1-%2-%3"))):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(ilvl))
        for tag, value in (("w:start", str(start)), ("w:numFmt", "decimal"), ("w:lvlText", text)):
            child = OxmlElement(tag)
            child.set(qn("w:val"), value)
            lvl.append(child)
        abstract.append(lvl)
    numbering.insert(0, abstract)

    used = [int(item.get(qn("w:numId"))) for item in numbering.findall(qn("w:num"))]
    num_id = str(max(used, default=0) + 1)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), num_id)
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), abstract_id)
    num.append(ref)
    if override_level1 is not None:
        override = OxmlElement("w:lvlOverride")
        override.set(qn("w:ilvl"), "1")
        start = OxmlElement("w:startOverride")
        start.set(qn("w:val"), str(override_level1))
        override.append(start)
        num.append(override)
    numbering.append(num)
    return num_id


def number_paragraph(paragraph, num_id: str, ilvl: int) -> None:
    numpr = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), str(ilvl))
    ref = OxmlElement("w:numId")
    ref.set(qn("w:val"), num_id)
    numpr.append(level)
    numpr.append(ref)
    paragraph._p.get_or_add_pPr().append(numpr)


def test_unseen_parent_levels_show_their_start_values():
    doc = Document()
    num_id = add_multilevel_list(doc)
    first = doc.add_paragraph("理解艺术设计的基本概念")
    second = doc.add_paragraph("掌握造型观察")
    number_paragraph(first, num_id, 2)
    number_paragraph(second, num_id, 2)

    labels = paragraph_numbers(doc)

    assert labels[first._p] == "1-1-1"
    assert labels[second._p] == "1-1-2"


def test_start_override_and_level_reset():
    """A parent level shown only implicitly does not consume a number: the
    first explicit level-1 item still shows the start value, the next one
    advances it, and each resets the level below."""
    doc = Document()
    num_id = add_multilevel_list(doc, override_level1=3)
    a = doc.add_paragraph("a")
    b = doc.add_paragraph("b")
    first_parent = doc.add_paragraph("group")
    c = doc.add_paragraph("c")
    second_parent = doc.add_paragraph("next group")
    d = doc.add_paragraph("d")
    for paragraph, level in ((a, 2), (b, 2), (first_parent, 1), (c, 2), (second_parent, 1), (d, 2)):
        number_paragraph(paragraph, num_id, level)

    labels = paragraph_numbers(doc)

    assert labels[a._p] == "1-3-1"
    assert labels[b._p] == "1-3-2"
    assert labels[first_parent._p] == "1-3"
    assert labels[c._p] == "1-3-1"
    assert labels[second_parent._p] == "1-4"
    assert labels[d._p] == "1-4-1"


def test_plain_paragraphs_get_no_label():
    doc = Document()
    plain = doc.add_paragraph("1-2-1 typed by hand")
    assert plain._p not in paragraph_numbers(doc)


def test_cell_text_prefixes_each_numbered_paragraph():
    doc = Document()
    num_id = add_multilevel_list(doc)
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.paragraphs[0].text = "理解艺术设计的基本概念"
    number_paragraph(cell.paragraphs[0], num_id, 2)
    number_paragraph(cell.add_paragraph("掌握造型观察"), num_id, 2)

    assert cell_text(cell, paragraph_numbers(doc)) == "1-1-1 理解艺术设计的基本概念 1-1-2 掌握造型观察"


def test_talent_plan_reads_auto_numbered_indicators(tmp_path):
    """The 2026 艺术设计 plan numbers groups 1-1 and 1-3 with Word lists and
    types the codes of 1-2 by hand; every group must come out the same."""
    doc = Document()
    list_11 = add_multilevel_list(doc)
    list_13 = add_multilevel_list(doc, override_level1=3)
    table = doc.add_table(rows=4, cols=4)
    table.cell(0, 1).text = "培养规格代码"
    table.cell(0, 2).text = "TOP10"
    table.cell(0, 3).text = "其他"
    table.cell(1, 0).text = "知识点"
    table.cell(1, 1).text = "1-1 艺术设计基础知识"
    table.cell(1, 2).paragraphs[0].text = "理解艺术设计的基本概念"
    number_paragraph(table.cell(1, 2).paragraphs[0], list_11, 2)
    number_paragraph(table.cell(1, 2).add_paragraph("掌握造型观察"), list_11, 2)
    table.cell(2, 0).text = "知识点"
    table.cell(2, 1).text = "1-2 文化内容与设计转化知识"
    table.cell(2, 2).text = "1-2-1 理解中华美学精神"
    table.cell(3, 0).text = "知识点"
    table.cell(3, 1).text = "1-3 用户研究与体验设计知识"
    table.cell(3, 2).paragraphs[0].text = "理解用户研究的作用"
    number_paragraph(table.cell(3, 2).paragraphs[0], list_13, 2)
    path = tmp_path / "auto-numbered.docx"
    doc.save(path)

    result = parse_talent_plan(path)

    assert [item.code for item in result.indicators] == ["1-1-1", "1-1-2", "1-2-1", "1-3-1"]
    assert result.indicators[0].description == "理解艺术设计的基本概念"
    assert result.indicators[3].description == "理解用户研究的作用"
