from collections.abc import Mapping, Sequence
from copy import deepcopy
from dataclasses import dataclass
import re
from typing import Any

from docx.oxml import OxmlElement
from docx.oxml.ns import qn


INFO_MARKERS = ("本次课标题", "授课学时", "本次课教学目标", "能力指标代码")
PROCESS_MARKERS = ("教学环节", "教学内容", "教师活动", "学生活动", "对应教学目标")
GENERATED_MARK_COLOR = "C00000"


class LessonTemplateError(ValueError):
    pass


@dataclass(frozen=True)
class LessonTemplateBlock:
    info_table_index: int
    process_table_index: int


@dataclass(frozen=True)
class LessonProcessStep:
    title: str
    minutes: int
    teacher_activity: str
    student_activity: str
    assessment: str


def _normalize(value: str) -> str:
    return "".join(value.split()).replace("（", "").replace("）", "").replace("(", "").replace(")", "")


def _table_text(table) -> str:
    return " ".join(_normalize(cell.text) for row in table.rows for cell in row.cells)


def _is_info_table(table) -> bool:
    text = _table_text(table)
    return all(marker in text for marker in INFO_MARKERS)


def _is_process_table(table) -> bool:
    text = " ".join(_normalize(cell.text) for row in table.rows[:2] for cell in row.cells)
    return all(marker in text for marker in PROCESS_MARKERS)


def find_lesson_template_blocks(document) -> list[LessonTemplateBlock]:
    blocks: list[LessonTemplateBlock] = []
    info_indexes = [index for index, table in enumerate(document.tables) if _is_info_table(table)]
    for info_position, info_index in enumerate(info_indexes):
        stop = info_indexes[info_position + 1] if info_position + 1 < len(info_indexes) else len(document.tables)
        process_index = next(
            (index for index in range(info_index + 1, stop) if _is_process_table(document.tables[index])),
            None,
        )
        if process_index is not None:
            blocks.append(LessonTemplateBlock(info_table_index=info_index, process_table_index=process_index))
    return blocks


def _get(item: Any, key: str, default: Any = "") -> Any:
    if isinstance(item, Mapping):
        return item.get(key, default)
    return getattr(item, key, default)


def set_cell_text(cell, value: Any, donor=None) -> None:
    text = str(value or "")
    paragraphs = cell.paragraphs
    for paragraph in paragraphs:
        for run in paragraph.runs:
            run.text = ""
    target = paragraphs[0]
    for paragraph in paragraphs[1:]:
        cell._tc.remove(paragraph._element)
    if target.runs:
        target.runs[0].text = text
        run = target.runs[0]
    else:
        run = target.add_run(text)
    # Blank template cells carry no run formatting, and the template's own
    # paragraph marks are not consistent either, so written text fell back to
    # the document default and came out larger than its neighbours.
    if donor is not None:
        existing = run._r.find(qn("w:rPr"))
        if existing is not None:
            run._r.remove(existing)
        run._r.insert(0, without_colour(donor))
        return
    _strip_colour(run)
    if run._r.find(qn("w:rPr")) is not None:
        return
    paragraph_properties = target._p.find(qn("w:pPr"))
    mark_properties = None if paragraph_properties is None else paragraph_properties.find(qn("w:rPr"))
    if mark_properties is not None:
        run._r.insert(0, deepcopy(mark_properties))


def without_colour(properties):
    """The red in the template only marked what to generate; output is black."""
    copied = deepcopy(properties)
    for tag in ("w:color", "w:highlight"):
        found = copied.find(qn(tag))
        if found is not None:
            copied.remove(found)
    return copied


def _strip_colour(run) -> None:
    properties = run._r.find(qn("w:rPr"))
    if properties is None:
        return
    for tag in ("w:color", "w:highlight"):
        found = properties.find(qn(tag))
        if found is not None:
            properties.remove(found)


def dominant_run_properties(table):
    """The formatting this table already uses for content, not for its labels.

    The school marks every cell it expects generated in dark red, so those runs
    are the template's own statement of how generated text should look -- 宋体
    body text, never the 黑体 of the label beside it. Where nothing is marked,
    fall back to whatever the table says most.

    Weigh by characters rather than by runs: Word splits "授课班级" into four
    one-character 黑体 runs, so counting runs made the labels outvote the
    content and the whole information table came out in 黑体.
    """
    marked = _most_used(_run_property_tally(table, marked_only=True))
    if marked is not None:
        return marked
    # A template with nothing marked and its content cells blank -- the school's
    # own 教学实施过程 table -- has only labels left to vote: 黑体 bold headers
    # across the top and 课前/课中/课后 down the side. The header alone outweighed
    # the one 宋体 小五 「课中」 cell (the rest of that column is merged into it)
    # and every generated row came out as a heading. Labels live in the header
    # row and the first column and announce themselves in bold, so the
    # regular-weight runs outside those places vote first, and the vote widens
    # only when they are silent.
    for skip_header, skip_first_column in ((True, True), (True, False), (False, False)):
        for regular_weight_only in (True, False):
            found = _most_used(
                _run_property_tally(
                    table,
                    marked_only=False,
                    skip_header=skip_header,
                    skip_first_column=skip_first_column,
                    regular_weight_only=regular_weight_only,
                )
            )
            if found is not None:
                return found
    return None


def _run_property_tally(
    table,
    marked_only: bool,
    skip_header: bool = False,
    skip_first_column: bool = False,
    regular_weight_only: bool = False,
) -> dict[tuple, dict[str, tuple[int, Any]]]:
    # Walk the runs themselves rather than row.cells: a merged cell is repeated
    # once per grid column it covers and would be counted that many times.
    tally: dict[tuple, dict[str, tuple[int, Any]]] = {}
    for run in _table_runs(table, skip_header, skip_first_column):
        properties = run.find(qn("w:rPr"))
        if properties is None:
            continue
        if marked_only and not _marks_generated_cell(properties):
            continue
        if regular_weight_only and _toggled(properties, "w:b"):
            continue
        text = "".join(node.text or "" for node in run.findall(qn("w:t"))).strip()
        if not text:
            continue
        colourless = without_colour(properties)
        variants = tally.setdefault(_visible_format(colourless), {})
        key = str(colourless.xml)
        weight, sample = variants.get(key, (0, colourless))
        variants[key] = (weight + len(text), sample)
    return tally


def _table_runs(table, skip_header: bool, skip_first_column: bool):
    for row_index, row in enumerate(table._tbl.findall(qn("w:tr"))):
        if skip_header and row_index == 0:
            continue
        for cell_index, cell in enumerate(row.findall(qn("w:tc"))):
            if skip_first_column and cell_index == 0:
                continue
            yield from cell.iter(qn("w:r"))


def _visible_format(properties) -> tuple:
    """What a reader actually sees: face, size and emphasis.

    Word sprinkles otherwise identical runs with w:lang, w:szCs and w:bCs, none
    of which change the page. Grouping on the raw XML split one font into five
    buckets and handed the vote to a minority formatting.
    """
    fonts = properties.find(qn("w:rFonts"))
    size = properties.find(qn("w:sz"))
    return (
        *(None if fonts is None else fonts.get(qn(attribute)) for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia")),
        None if size is None else size.get(qn("w:val")),
        *(_toggled(properties, tag) for tag in ("w:b", "w:i", "w:u", "w:strike")),
    )


def _toggled(properties, tag: str) -> bool:
    element = properties.find(qn(tag))
    return element is not None and element.get(qn("w:val")) not in ("0", "false", "none")


def _marks_generated_cell(properties) -> bool:
    colour = properties.find(qn("w:color"))
    return colour is not None and colour.get(qn("w:val")) == GENERATED_MARK_COLOR


def _most_used(tally: dict[tuple, dict[str, tuple[int, Any]]]):
    if not tally:
        return None
    heaviest = max(tally.values(), key=lambda variants: sum(weight for weight, _ in variants.values()))
    return max(heaviest.values(), key=lambda item: item[0])[1]


SECTION_HEADING = "教学基本情况"
# The page title sits above the first lesson table and reads "<school name>教案".
# Matching a literal school name would tie the export to one school, so the line
# is recognised by its shape: a short heading that names the document itself.
PAGE_TITLE_PATTERN = re.compile(r"^.{0,30}教\s*案.{0,10}$")


def move_section_heading_above_first_block(document) -> bool:
    """The school's template has "一、教学基本情况" sitting under its table."""
    blocks = find_lesson_template_blocks(document)
    if not blocks:
        return False
    first_table = document.tables[blocks[0].info_table_index]._tbl
    body = document.element.body
    for paragraph in document.paragraphs:
        if SECTION_HEADING not in _normalize(paragraph.text):
            continue
        if list(body).index(paragraph._p) < list(body).index(first_table):
            return False
        first_table.addprevious(paragraph._p)
        return True
    return False


def _move_table_into_text_flow(table) -> None:
    positioning = table._tbl.tblPr.find(qn("w:tblpPr"))
    if positioning is not None:
        table._tbl.tblPr.remove(positioning)


def _find_label(table, label: str) -> tuple[int, int] | None:
    normalized = _normalize(label)
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            if normalized in _normalize(cell.text):
                return row_index, column_index
    return None


def _set_right_of_label(table, label: str, value: Any, donor=None) -> bool:
    location = _find_label(table, label)
    if location is None:
        return False
    row_index, column_index = location
    label_cell = table.cell(row_index, column_index)
    for target_index in range(column_index + 1, len(table.rows[row_index].cells)):
        target = table.cell(row_index, target_index)
        if target._tc is not label_cell._tc:
            set_cell_text(target, value, donor)
            return True
    return False


def _set_below_label(table, label: str, value: Any, donor=None) -> bool:
    location = _find_label(table, label)
    if location is None or location[0] + 1 >= len(table.rows):
        return False
    set_cell_text(table.cell(location[0] + 1, location[1]), value, donor)
    return True


def _lesson_hours(lesson: Any) -> str:
    minutes = int(_get(lesson, "duration_minutes", 0) or 0)
    return f"{minutes / 40:g}" if minutes else ""


def _class_time(context: Any) -> str:
    week = _get(context, "week_no")
    weekday = _get(context, "weekday")
    periods = _get(context, "periods")
    if not any((week, weekday, periods)):
        return ""
    weekday_text = str(weekday or "")
    if weekday_text and not weekday_text.startswith("周"):
        weekday_text = f"周{weekday_text}"
    parts = [f"第 {week} 周" if week else "", weekday_text, f"第 {periods} 节" if periods else ""]
    return " ".join(part for part in parts if part)


def _fill_info_table(table, lesson: Any, context: Any, common_values: Mapping[str, Any]) -> None:
    donor = dominant_run_properties(table)
    values = {
        "课程名称": common_values.get("课程名称", ""),
        "任课教师": common_values.get("任课教师", ""),
        "本次课标题": _get(lesson, "title"),
        "授课学时": _lesson_hours(lesson),
        "授课班级": _get(context, "class_name") or common_values.get("授课班级", ""),
        "上课时间": _class_time(context),
        "上课地点": _get(context, "location") or common_values.get("上课地点", ""),
    }
    # These four rows are red in the school's template, meaning the teacher
    # expects them generated. Leaving them alone kept another lesson's text.
    values.update(
        {
            "教学重点": _get(lesson, "key_points"),
            "教学难点": _get(lesson, "difficult_points"),
            "教学方法手段": _get(context, "teaching_methods"),
            "教学资源": _get(lesson, "teaching_preparation"),
        }
    )
    for label, value in values.items():
        if value != "":
            _set_right_of_label(table, label, value, donor)

    _fill_objective_rows(table, lesson, donor)
    _clear_unused_session_rows(table)
    _set_below_label(table, "教学准备", _get(lesson, "teaching_preparation"), donor)
    _set_below_label(table, "课堂小结", _get(lesson, "summary"), donor)


def _clear_unused_session_rows(table) -> None:
    """The template stacks blank 授课班级 rows for extra teaching classes."""
    found = _find_label(table, "授课班级")
    if found is None:
        return
    label_row, label_column = found
    for row_index in range(label_row + 1, len(table.rows)):
        row = table.rows[row_index]
        if _normalize(row.cells[label_column].text) != _normalize(table.rows[label_row].cells[label_column].text):
            break
        for cell in row.cells:
            if _normalize(cell.text) == _normalize(table.rows[label_row].cells[label_column].text):
                continue
            if re.fullmatch(r"[第周节\s\-－]*", cell.text.strip() or "-") or "第" in cell.text:
                set_cell_text(cell, "")


def _goal_statements(value: Any) -> list[str]:
    lines = [line.strip() for line in str(value or "").splitlines() if line.strip()]
    return [
        re.sub(r"^(?:K\d+|\d+)[\.、:：)）]\s*", "", line).strip()
        for line in lines
    ]


def _fill_objective_rows(table, lesson: Any, donor=None) -> None:
    header = _find_label(table, "本次课教学目标")
    goal_header = _find_label(table, "课程教学目标")
    ability_header = _find_label(table, "能力指标代码")
    if header is None:
        return
    start_row = header[0] + 1
    stop_row = len(table.rows)
    for row_index in range(start_row, len(table.rows)):
        row_text = _normalize(" ".join(cell.text for cell in table.rows[row_index].cells))
        if any(marker in row_text for marker in ("教学重点", "教学难点", "教学方法", "教学资源")):
            stop_row = row_index
            break

    statements = _goal_statements(_get(lesson, "teaching_goals"))
    goal_column = goal_header[1] if goal_header is not None else None
    ability_column = ability_header[1] if ability_header is not None else None
    # Codes line up with objectives one for one when the model returns them per
    # goal; a single shared line still applies to every row, as before.
    goal_codes = _per_goal_codes(_get(lesson, "course_goal_codes"), len(statements))
    ability_codes = _per_goal_codes(_get(lesson, "ability_codes"), len(statements))
    for position, row_index in enumerate(range(start_row, stop_row)):
        statement = statements[position] if position < len(statements) else ""
        set_cell_text(table.cell(row_index, 0), f"K{position + 1}：{statement}" if statement else "", donor)
        if goal_column is not None:
            set_cell_text(table.cell(row_index, goal_column), goal_codes[position] if statement else "", donor)
        if ability_column is not None:
            set_cell_text(table.cell(row_index, ability_column), ability_codes[position] if statement else "", donor)


def _per_goal_codes(value: Any, count: int) -> list[str]:
    lines = [line.strip() for line in str(value or "").splitlines() if line.strip()]
    if len(lines) >= count > 0:
        return lines[:count]
    shared = " ".join(lines)
    return [shared] * max(count, 0)


def _split_process(value: str) -> dict[str, str]:
    result: dict[str, str] = {}
    for raw_line in str(value or "").splitlines():
        line = raw_line.strip()
        if not line:
            continue
        separator = "：" if "：" in line else ":" if ":" in line else None
        if separator is None:
            result.setdefault("其他", line)
            continue
        label, content = line.split(separator, 1)
        result[_normalize(label)] = content.strip()
    return result


def _process_columns(table) -> dict[str, int]:
    columns: dict[str, int] = {}
    for index, cell in enumerate(table.rows[0].cells):
        header = _normalize(cell.text)
        for marker in ("教学内容", "教师活动", "学生活动", "对应教学目标"):
            if marker in header and marker not in columns:
                columns[marker] = index
        if "时长" in header and "教学环节" not in header and "时长" not in columns:
            columns["时长"] = index
    if "时长" not in columns and columns.get("教学内容", 0) > 1:
        columns["时长"] = columns["教学内容"] - 1
    return columns


def _parse_process_steps(value: Any) -> list[LessonProcessStep]:
    steps: list[LessonProcessStep] = []
    current: dict[str, Any] | None = None
    title_pattern = re.compile(r"^(.+?)[（(]\s*(\d+)\s*分钟\s*[）)]$")
    for raw_line in str(value or "").splitlines():
        line = raw_line.strip()
        if not line:
            continue
        title_match = title_pattern.fullmatch(line)
        if title_match:
            if current is not None:
                steps.append(_process_step(current))
            current = {"title": title_match.group(1).strip(), "minutes": int(title_match.group(2))}
            continue
        if current is None:
            continue
        separator = "：" if "：" in line else ":" if ":" in line else None
        if separator is None:
            continue
        label, content = line.split(separator, 1)
        normalized = _normalize(label)
        if normalized == "教师活动":
            current["teacher_activity"] = content.strip()
        elif normalized == "学生活动":
            current["student_activity"] = content.strip()
        elif normalized in {"学习评价", "评价"}:
            current["assessment"] = content.strip()
    if current is not None:
        steps.append(_process_step(current))
    return steps


def _process_step(values: Mapping[str, Any]) -> LessonProcessStep:
    return LessonProcessStep(
        title=str(values.get("title", "")),
        minutes=int(values.get("minutes", 0) or 0),
        teacher_activity=str(values.get("teacher_activity", "")),
        student_activity=str(values.get("student_activity", "")),
        assessment=str(values.get("assessment", "")),
    )


def _legacy_process_steps(process: Mapping[str, str]) -> list[LessonProcessStep]:
    return [
        LessonProcessStep(
            title=label,
            minutes=0,
            teacher_activity=f"组织并指导：{content}",
            student_activity=f"参与并完成：{content}",
            assessment="检查课堂任务完成情况。",
        )
        for label, content in process.items()
        if label not in {"课前", "课后"}
    ]


def _merge_process_steps(steps: Sequence[LessonProcessStep]) -> LessonProcessStep:
    return LessonProcessStep(
        title="\n".join(step.title for step in steps if step.title),
        minutes=sum(step.minutes for step in steps),
        teacher_activity="\n".join(step.teacher_activity for step in steps if step.teacher_activity),
        student_activity="\n".join(step.student_activity for step in steps if step.student_activity),
        assessment="\n".join(step.assessment for step in steps if step.assessment),
    )


def _fill_process_table(table, lesson: Any, context: Any = None) -> None:
    donor = dominant_run_properties(table)
    columns = _process_columns(table)
    lesson_goal_codes = [f"K{index}" for index, _ in enumerate(_goal_statements(_get(lesson, "teaching_goals")), start=1)]
    process = _split_process(_get(lesson, "teaching_process"))
    post_content = _get(lesson, "homework") or process.get("课后", "")
    steps = _parse_process_steps(_get(lesson, "teaching_process")) or _legacy_process_steps(process)
    pre_rows: list[Any] = []
    in_rows: list[Any] = []
    post_rows: list[Any] = []
    for row in table.rows[1:]:
        phase = _normalize(row.cells[0].text)
        if "课前" in phase:
            pre_rows.append(row)
        elif "课后" in phase:
            post_rows.append(row)
        else:
            in_rows.append(row)

    row_values: list[tuple[Any, LessonProcessStep | None, str, str, str]] = []
    preparation = _get(lesson, "teaching_preparation") or process.get("课前", "")
    pre_content = process.get("课前") or ("教学准备" if preparation else "")
    row_values.extend(
        (
            row,
            None,
            pre_content if index == 0 else "",
            "发布课前学习任务并检查学习资源。" if index == 0 else "",
            # The student column is what students do, not the equipment list;
            # that belongs in 教学资源 on the information table.
            (_get(context, "pre_task") or "预习本次课内容，做好上课准备。") if index == 0 else "",
        )
        for index, row in enumerate(pre_rows)
    )
    for index, row in enumerate(in_rows):
        if index == len(in_rows) - 1 and len(steps) > len(in_rows):
            step = _merge_process_steps(steps[index:])
        else:
            step = steps[index] if index < len(steps) else None
        content = ""
        if step is not None:
            content = step.title
            if step.assessment:
                content += f"\n学习评价：{step.assessment}"
        row_values.append(
            (
                row,
                step,
                content,
                step.teacher_activity if step is not None else "",
                step.student_activity if step is not None else "",
            )
        )
    row_values.extend(
        (
            row,
            None,
            ("课后巩固与提交" if post_content else "") if index == 0 else "",
            post_content if index == 0 else "",
            # Distinct from the assignment text so the row is not the same
            # sentence three times.
            (_get(context, "post_task") or post_content) if index == 0 else "",
        )
        for index, row in enumerate(post_rows)
    )

    for row, step, content, teacher_activity, student_activity in row_values:
        set_cell_text(row.cells[columns["教学内容"]], content, donor)
        set_cell_text(row.cells[columns["教师活动"]], teacher_activity, donor)
        set_cell_text(row.cells[columns["学生活动"]], student_activity, donor)
        # This column names the lesson's own K objectives, not the course-level
        # M codes; the template's own example reads "K1", "K2 K3".
        set_cell_text(row.cells[columns["对应教学目标"]], " ".join(lesson_goal_codes) if content else "", donor)
        if "时长" in columns:
            duration_cell = row.cells[columns["时长"]]
            # On 课前 and 课后 rows the template merges 教学环节 with 时长, so
            # writing a blank duration erased the phase label itself.
            if duration_cell._tc is not row.cells[0]._tc:
                set_cell_text(duration_cell, f"{step.minutes}分钟" if step is not None and step.minutes else "", donor)


def read_process_layout(document) -> list[int]:
    """Minutes the template fixes for each 课中 row, empty when it leaves them blank.

    The school's template writes the pacing itself and colours only the cells
    the teacher expects AI to fill, so generation should match this shape
    instead of inventing its own segment count.
    """
    blocks = find_lesson_template_blocks(document)
    if not blocks:
        return []
    table = document.tables[blocks[0].process_table_index]
    columns = _process_columns(table)
    if "时长" not in columns:
        return []
    minutes: list[int] = []
    for row in table.rows[1:]:
        phase = _normalize(row.cells[0].text)
        if "课前" in phase or "课后" in phase:
            continue
        found = re.findall(r"\d+", row.cells[columns["时长"]].text)
        if not found:
            return []
        minutes.append(int(found[0]))
    return minutes


def _repeat_last_block(document, blocks: list[LessonTemplateBlock], wanted: int) -> None:
    """Clone the template's final lesson block until there is one per session.

    Schools hand out a template holding a single blank lesson, so the export
    has to grow it rather than demand one block per session up front.
    """
    body = document.element.body
    children = list(body)
    last = blocks[-1]
    start = children.index(document.tables[last.info_table_index]._tbl)
    # Walk back over the centred title and the 一、/二、 headings so every
    # cloned session carries the same page furniture as the first one.
    while start > 0:
        previous = children[start - 1]
        if not previous.tag.endswith("}p"):
            break
        text = "".join(previous.itertext()).strip()
        if not text or PAGE_TITLE_PATTERN.match(text) or re.match(r"^[一二三四五六七八九十]、", text):
            start -= 1
            continue
        break
    end = len(children)
    for index in range(len(children) - 1, start, -1):
        tag = children[index].tag
        if tag.endswith("}sectPr"):
            end = index
    source = children[start:end]
    if not source:
        return

    anchor = source[-1]
    for _ in range(wanted - len(blocks)):
        page_break = OxmlElement("w:p")
        run = OxmlElement("w:r")
        broken = OxmlElement("w:br")
        broken.set(qn("w:type"), "page")
        run.append(broken)
        page_break.append(run)
        anchor.addnext(page_break)
        anchor = page_break
        for element in source:
            copied = deepcopy(element)
            anchor.addnext(copied)
            anchor = copied


def fill_structured_lesson_template(
    document,
    lessons: Sequence[Any],
    contexts: Sequence[Any],
    common_values: Mapping[str, Any],
) -> None:
    move_section_heading_above_first_block(document)
    blocks = find_lesson_template_blocks(document)
    if not blocks:
        raise LessonTemplateError("教案模板未识别，请添加占位符或使用学校标准模板")
    if len(blocks) < len(lessons):
        _repeat_last_block(document, blocks, len(lessons))
        blocks = find_lesson_template_blocks(document)
    if len(blocks) < len(lessons):
        raise LessonTemplateError(f"教案模板仅识别到 {len(blocks)} 个课次区块，需要 {len(lessons)} 个")

    for index, lesson in enumerate(lessons):
        block = blocks[index]
        context = contexts[index] if index < len(contexts) else {}
        info_table = document.tables[block.info_table_index]
        process_table = document.tables[block.process_table_index]
        _move_table_into_text_flow(info_table)
        _move_table_into_text_flow(process_table)
        _fill_info_table(info_table, lesson, context, common_values)
        _fill_process_table(process_table, lesson, context)


# Some red cells are filled from the course record, not by the model: when the
# template was written for the same course they legitimately match, so matching
# text there is not evidence of a gap.
DATA_DRIVEN_LABELS = ("课程名称", "任课教师", "适用专业", "课程类型", "授课班级", "上课时间", "上课地点")


def _cell_marked_runs(cell) -> list[str]:
    marked = []
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            colour = run.font.color
            if colour is not None and colour.rgb is not None and str(colour.rgb) == GENERATED_MARK_COLOR:
                text = run.text.strip()
                if text:
                    marked.append(text)
    return marked


def find_unfilled_generated_cells(template, produced) -> list[str]:
    """Cells the template marks red but the export left at their template text.

    The school colours every cell it expects AI to fill. Comparing the produced
    document back against that marking catches a whole class of silent defects
    -- a row nobody wired up keeps another lesson's content and still looks
    like a finished document.
    """
    template_blocks = find_lesson_template_blocks(template)
    produced_blocks = find_lesson_template_blocks(produced)
    if not template_blocks or not produced_blocks:
        return []

    reference = template.tables[template_blocks[0].info_table_index]
    marked_cells: list[tuple[int, int, str]] = []
    for row_index, row in enumerate(reference.rows):
        for column_index, cell in enumerate(row.cells):
            marked = _cell_marked_runs(cell)
            if marked:
                marked_cells.append((row_index, column_index, _normalize(cell.text)))

    stale: list[str] = []
    for position, block in enumerate(produced_blocks, start=1):
        table = produced.tables[block.info_table_index]
        seen: set[str] = set()
        for row_index, column_index, template_text in marked_cells:
            if row_index >= len(table.rows) or column_index >= len(table.rows[row_index].cells):
                continue
            actual = _normalize(table.rows[row_index].cells[column_index].text)
            if not template_text or actual != template_text:
                continue
            row_text = _normalize(" ".join(cell.text for cell in table.rows[row_index].cells))
            if any(marker in row_text for marker in DATA_DRIVEN_LABELS):
                continue
            # A bare code such as M1 or 010101A1 says nothing on its own.
            if re.fullmatch(r"[A-Za-z0-9\-]{1,12}", template_text):
                continue
            label = _normalize(table.rows[row_index].cells[0].text)[:12] or f"第{row_index}行"
            if label in seen:
                continue
            seen.add(label)
            stale.append(f"第 {position} 次课「{label}」仍是模板原文")
    return stale
