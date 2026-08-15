from pathlib import Path

from docx import Document


MATERIAL_TYPE_LABELS = {"assignment": "作业", "test": "测试"}
DIFFICULTY_LABELS = {"basic": "基础", "medium": "适中", "advanced": "提高"}


def export_session_material_docx(
    output_path: Path | str,
    task,
    outline,
    material,
) -> None:
    document = Document()
    document.add_heading(material.title, level=1)

    table = document.add_table(rows=4, cols=4)
    details = (
        ("课程名称", task.course_name, "授课班级", task.class_name),
        ("任课教师", task.teacher_name, "课次", f"第 {outline.session_no} 次课"),
        ("教学主题", outline.topic, "上课时间", f"{outline.date_text} {outline.weekday} 第 {outline.periods} 节"),
        (
            "材料信息",
            MATERIAL_TYPE_LABELS.get(material.material_type, material.material_type),
            "难度与用时",
            f"{DIFFICULTY_LABELS.get(material.difficulty, material.difficulty)} / {material.estimated_minutes} 分钟",
        ),
    )
    for row, values in zip(table.rows, details, strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = str(value)

    _add_section(document, "材料内容", material.content)
    _add_section(document, "参考答案", material.reference_answer)
    _add_section(document, "评分标准", material.grading_criteria)
    _add_section(
        document,
        "关联依据",
        f"课程目标：{material.course_goal_codes or '未填写'}\n能力指标：{material.ability_codes or '未填写'}",
    )
    document.save(str(output_path))


def _add_section(document, title: str, content: str) -> None:
    document.add_heading(title, level=2)
    document.add_paragraph(content or "未填写")
