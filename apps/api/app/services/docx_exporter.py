from copy import deepcopy
from pathlib import Path
import re
from typing import Any, Iterable, Mapping

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.services.lesson_template_filler import (
    LessonTemplateError,
    dominant_run_properties,
    fill_structured_lesson_template,
    find_unfilled_generated_cells,
    set_cell_text,
)
from app.services.outline_template_filler import (
    OutlineTemplateError,
    clear_generated_marking,
    fill_outline_sections,
    find_unfilled_generated_sections,
    template_marks_generated_sections,
)


HEADER_FIELDS = (
    ("课程思政", "ideological_point"),
    ("教学方法", "teaching_methods"),
    ("课程目标", "course_goal_codes"),
    ("能力指标", "ability_codes"),
    ("教学主题", "topic"),
    ("主题", "topic"),
    ("教学内容", "teaching_content"),
    ("关键教学", "in_class_task"),
    ("教学要求", "in_class_task"),
    ("学习要求", "learning_tasks"),
    ("学习任务", "learning_tasks"),
    ("日期", "date_text"),
    ("周次", "week_no"),
    ("星期", "weekday"),
    ("节次", "periods"),
    ("课次", "session_no"),
    ("备注", "note"),
)
REQUIRED_SCHEDULE_FIELDS = {"date_text", "week_no", "periods", "teaching_content"}
# How the prose columns of 学习进程 divide the width they already occupy. 日期,
# 周次 and 节次 hold fixed-format values and keep the width the template drew.
SCHEDULE_COLUMN_SHARES = {
    "teaching_content": 30,
    "ideological_point": 21,
    "teaching_methods": 21,
    "in_class_task": 28,
    "learning_tasks": 28,
}


COURSE_NAME_LABEL = "课程名称"
_COURSE_DECLARATION = re.compile(rf"\s*{COURSE_NAME_LABEL}\s*[：:]\s*(\S.*)")


def _normalize_course(value: str) -> str:
    return "".join(value.split()).replace("《", "").replace("》", "")


def _declared_course_names(document) -> list[str]:
    """Courses the template names outright.

    A blank template -- headings and empty tables -- names none, and that is the
    intended way to start one. A template that names a course is somebody's
    finished outline: everything in it except 学习进程 ships to the reader
    untouched, so the course it names has to be this one.
    """
    found: list[str] = []

    def note(value: str) -> None:
        name = value.strip().splitlines()[0].strip() if value.strip() else ""
        if name and name not in found:
            found.append(name)

    for paragraph in document.paragraphs:
        matched = _COURSE_DECLARATION.match(paragraph.text)
        if matched:
            note(matched.group(1))

    for table in document.tables:
        for row in table.rows:
            cells = row.cells
            for index, cell in enumerate(cells):
                text = cell.text.strip()
                matched = _COURSE_DECLARATION.match(text)
                if matched:
                    note(matched.group(1))
                elif text == COURSE_NAME_LABEL and index + 1 < len(cells):
                    note(cells[index + 1].text)
    return found


def _conflicting_course_name(document, course_name: str) -> str:
    if not course_name.strip():
        return ""
    wanted = _normalize_course(course_name)
    for declared in _declared_course_names(document):
        if _normalize_course(declared) != wanted:
            return declared
    return ""


def _replace_in_paragraph(paragraph, values: dict[str, str]) -> None:
    original = paragraph.text
    replaced = original
    for key, value in values.items():
        replaced = replaced.replace(f"{{{{{key}}}}}", value)

    if replaced == original:
        return

    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = replaced
    else:
        paragraph.add_run(replaced)


def _replace_placeholders(document, values: dict[str, str]) -> None:
    for paragraph in document.paragraphs:
        _replace_in_paragraph(paragraph, values)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace_in_paragraph(paragraph, values)


def _all_document_text(document) -> str:
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _normalize_header(value: str) -> str:
    return "".join(value.split()).replace("、", "").replace("，", "")


def _schedule_columns(row) -> dict[int, str]:
    columns: dict[int, str] = {}
    for index, cell in enumerate(row.cells):
        header = _normalize_header(cell.text)
        for keyword, field in HEADER_FIELDS:
            if keyword in header:
                columns[index] = field
                break
    return columns


def _find_schedule_table(document):
    for table in document.tables:
        for row_index, row in enumerate(table.rows[:5]):
            columns = _schedule_columns(row)
            if REQUIRED_SCHEDULE_FIELDS.issubset(columns.values()):
                return table, row_index, columns
    return None


def _row_value(row: Any, field: str) -> str:
    if field == "learning_tasks":
        parts = (
            ("课前", _get_value(row, "pre_task")),
            ("课中", _get_value(row, "in_class_task")),
            ("课后", _get_value(row, "post_task")),
        )
        return "\n".join(f"{label}：{value}" for label, value in parts if value)
    return str(_get_value(row, field) or "")


def _get_value(row: Any, field: str) -> Any:
    if isinstance(row, Mapping):
        return row.get(field, "")
    return getattr(row, field, "")


def _add_bold_paragraph(document, text: str):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True
    return paragraph


def _ensure_data_rows(table, header_row_index: int, count: int) -> None:
    existing_count = len(table.rows) - header_row_index - 1
    while existing_count < count:
        if len(table.rows) > header_row_index + 1:
            table._tbl.append(deepcopy(table.rows[-1]._tr))
        else:
            table.add_row()
        existing_count += 1


def _fill_schedule_table(document, rows: Iterable[Any]) -> bool:
    schedule = _find_schedule_table(document)
    if schedule is None:
        return False

    table, header_row_index, columns = schedule
    donor = dominant_run_properties(table)
    widths = _schedule_column_widths(table, header_row_index, columns)
    outline_rows = list(rows)
    _ensure_data_rows(table, header_row_index, len(outline_rows))
    data_rows = table.rows[header_row_index + 1 :]

    for data_index, table_row in enumerate(data_rows):
        outline_row = outline_rows[data_index] if data_index < len(outline_rows) else None
        for column_index, field in columns.items():
            value = _row_value(outline_row, field) if outline_row is not None else ""
            set_cell_text(table_row.cells[column_index], value, donor)

    # Blanking only the mapped columns left the template's own text standing in
    # every other column of the surplus rows, so drop those rows outright.
    for table_row in data_rows[len(outline_rows) :]:
        table_row._tr.getparent().remove(table_row._tr)

    _apply_column_widths(table, widths)
    return True


def _cell_width(cell) -> int | None:
    width = cell._tc.tcPr.find(qn("w:tcW")) if cell._tc.tcPr is not None else None
    value = None if width is None else width.get(qn("w:w"))
    return None if value is None or not value.isdigit() else int(value)


def _schedule_column_widths(table, header_row_index: int, columns: dict[int, str]) -> list[int | None]:
    """Redivide the prose columns; leave the fixed-format ones as drawn.

    The template's grid was measured for entries like "讲授" and "课堂纪律要求",
    so 课程思政 got 709 twips -- about two characters. Generated rows carry a
    sentence in every column, and Word's autofit reflowed the table differently
    in each document. These shares are the ones the teacher settled on by hand.
    """
    widths = [_cell_width(cell) for cell in table.rows[header_row_index].cells]
    shares = {
        index: SCHEDULE_COLUMN_SHARES[field]
        for index, field in columns.items()
        if field in SCHEDULE_COLUMN_SHARES and index < len(widths) and widths[index] is not None
    }
    if len(shares) < 2:
        return widths

    available = sum(widths[index] for index in shares)
    total = sum(shares.values())
    for index, share in shares.items():
        widths[index] = round(available * share / total)
    # Rounding must not move the table's right edge.
    widest = max(shares, key=lambda index: widths[index])
    widths[widest] += available - sum(widths[index] for index in shares)
    return widths


def _apply_column_widths(table, widths: list[int | None]) -> None:
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for column, width in zip(grid.findall(qn("w:gridCol")), widths):
            if width is not None:
                column.set(qn("w:w"), str(width))
    layout = table._tbl.tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table._tbl.tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")
    for row in table.rows:
        # A horizontally merged cell repeats across the grid columns it covers,
        # and its width is their sum rather than the last one's.
        spans: list[tuple[Any, list[int]]] = []
        for index, cell in enumerate(row.cells):
            if spans and spans[-1][0]._tc is cell._tc:
                spans[-1][1].append(index)
            else:
                spans.append((cell, [index]))
        for cell, covered in spans:
            if any(index >= len(widths) or widths[index] is None for index in covered):
                continue
            properties = cell._tc.get_or_add_tcPr()
            width = properties.find(qn("w:tcW"))
            if width is None:
                width = OxmlElement("w:tcW")
                properties.append(width)
            width.set(qn("w:w"), str(sum(widths[index] for index in covered)))
            width.set(qn("w:type"), "dxa")


def fill_docx_placeholders(
    template_path: Path | str,
    output_path: Path | str,
    values: dict[str, str],
) -> None:
    document = Document(str(template_path))
    _replace_placeholders(document, values)
    document.save(str(output_path))


def fill_outline_docx(
    template_path: Path | str,
    output_path: Path | str,
    values: dict[str, str],
    rows: Iterable[Any],
    sections: Any = None,
    resources: Any = None,
    assessments: Any = None,
) -> None:
    document = Document(str(template_path))
    # After substitution, so a template written with {{课程名称}} resolves to this
    # course instead of looking like it belongs to nobody.
    _replace_placeholders(document, values)
    conflict = _conflicting_course_name(document, str(values.get("课程名称", "") or ""))
    if conflict:
        raise OutlineTemplateError(
            f"导出中止：模板里写的课程是《{conflict}》，本次任务是《{values.get('课程名称', '')}》。"
            "除「学习进程」表以外，模板正文会原样导出，"
            "照此导出会得到一份内容属于另一门课的大纲。"
            "请改用本课程的大纲模板，或使用只有标题和空表的空白模板。"
        )
    marked = template_marks_generated_sections(document)
    if marked and sections is None:
        raise OutlineTemplateError(
            "导出中止：这份模板把正文各节标成了「【AI 生成】」，但本次任务还没有生成大纲正文。"
            "请先生成大纲正文，或改用不需要生成正文的模板。"
        )
    _fill_schedule_table(document, rows)
    if marked:
        fill_outline_sections(document, sections, resources, assessments)
        # The template states which sections it expects generated; one still
        # showing its marker would ship an instruction to the teacher as if it
        # were the course's own text.
        stale = find_unfilled_generated_sections(document)
        if stale:
            raise OutlineTemplateError(
                "导出中止：以下内容仍是模板占位符，未被生成结果替换——"
                + "；".join(stale[:6])
                + (f"（共 {len(stale)} 处）" if len(stale) > 6 else "")
            )
        clear_generated_marking(document)

    document.save(str(output_path))


def fill_lesson_docx(
    template_path: Path | str,
    output_path: Path | str,
    values: dict[str, str],
    lessons: Iterable[Any],
    contexts: Iterable[Any] | None = None,
) -> None:
    lesson_list = list(lessons)
    document = Document(str(template_path))
    has_placeholder = "{{教案正文}}" in _all_document_text(document)
    if has_placeholder:
        lesson_text = values.get("教案正文") or _lesson_plans_to_text(lesson_list)
        _replace_placeholders(document, {**values, "教案正文": lesson_text})
    else:
        _replace_placeholders(document, values)
        fill_structured_lesson_template(document, lesson_list, list(contexts or []), values)
        # The template colours every cell it expects generated; a cell still
        # holding its template text means another lesson's content shipped.
        stale = find_unfilled_generated_cells(Document(str(template_path)), document)
        if stale:
            raise LessonTemplateError(
                "导出中止：以下内容仍是模板原文，未被本次课的生成结果替换——"
                + "；".join(stale[:6])
                + (f"（共 {len(stale)} 处）" if len(stale) > 6 else "")
            )

    document.save(str(output_path))


def _lesson_plans_to_text(lessons: list[Any]) -> str:
    return "\n\n".join(_lesson_to_text(lesson) for lesson in lessons)


def _lesson_to_text(lesson: Any) -> str:
    return "\n".join(
        [
            _row_value(lesson, "title"),
            f"授课时长：{_row_value(lesson, 'duration_minutes')} 分钟",
            f"教学目标：{_row_value(lesson, 'teaching_goals')}",
            f"教学重点：{_row_value(lesson, 'key_points')}",
            f"教学难点：{_row_value(lesson, 'difficult_points')}",
            f"教学准备：{_row_value(lesson, 'teaching_preparation')}",
            f"教学过程：\n{_row_value(lesson, 'teaching_process')}",
            f"课堂小结：{_row_value(lesson, 'summary')}",
            f"课后任务：{_row_value(lesson, 'homework')}",
            f"课程目标：{_row_value(lesson, 'course_goal_codes')}",
            f"能力指标：{_row_value(lesson, 'ability_codes')}",
            f"教学反思：{_row_value(lesson, 'reflection')}",
        ]
    )
