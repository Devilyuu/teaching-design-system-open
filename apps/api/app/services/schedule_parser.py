import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


@dataclass(frozen=True)
class ScheduleSession:
    week_no: int
    date_text: str
    weekday: str
    periods: str
    course_name: str
    class_name: str
    location: str
    hours: int


@dataclass(frozen=True)
class ColumnMatch:
    field: str
    column_index: int
    header_text: str
    confidence: str


@dataclass(frozen=True)
class ScheduleAnalysis:
    sessions: list[ScheduleSession]
    header_row_index: int
    matches: list[ColumnMatch]
    detected_headers: list[str] = field(default_factory=list)
    unmapped_headers: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    course_names: list[str] = field(default_factory=list)
    teaching_classes: list[str] = field(default_factory=list)


class ScheduleParseError(ValueError):
    def __init__(self, message: str, detected_headers: list[str], missing_fields: list[str]) -> None:
        super().__init__(message)
        self.detected_headers = detected_headers
        self.missing_fields = missing_fields


FIELD_ORDER = (
    "week_no",
    "date_text",
    "weekday",
    "periods",
    "course_name",
    "class_name",
    "location",
)

FIELD_LABELS = {
    "week_no": "周次",
    "date_text": "日期",
    "weekday": "星期",
    "periods": "节次",
    "course_name": "课程",
    "class_name": "班级",
    "location": "地点",
}

# Longest aliases first so 课程名称 wins over 课程 when both could match.
FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "week_no": ("上课周次", "教学周次", "教学周", "周次", "周数", "周"),
    "date_text": ("上课日期", "授课日期", "日期"),
    "weekday": ("星期几", "星期", "周几", "礼拜"),
    "periods": ("上课节次", "节次", "课节", "节数", "节"),
    "course_name": ("课程名称", "课程名", "课程", "科目"),
    "class_name": ("授课班级", "上课班级", "教学班级", "班级名称", "教学班", "班级"),
    "location": ("上课地点", "上课教室", "上课地址", "教室", "地点", "场地"),
}

# Columns whose header carries one of these never hold the value we want, even
# when the header also contains a field alias (课程代码 is not 课程).
EXCLUDE_HINTS = ("代码", "编号", "序号", "学分", "性质", "类别", "类型", "备注", "总学时", "周学时")

REQUIRED_FIELDS = ("periods",)

HEADER_SCAN_ROWS = 15
MIN_HEADER_MATCHES = 3

WEEKDAY_CHARS = ("一", "二", "三", "四", "五", "六", "日")

_FULLWIDTH = {code: code - 0xFEE0 for code in range(0xFF01, 0xFF5F)}
_FULLWIDTH[0x3000] = 0x20


def _norm(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    text = str(value).translate(_FULLWIDTH).strip()
    return re.sub(r"\s+", " ", text)


def _norm_header(value) -> str:
    return _norm(value).replace(" ", "").rstrip(":：")


def _normalize_periods(raw: str) -> tuple[str, int]:
    numbers = [int(item) for item in re.findall(r"\d+", raw)]
    if not numbers:
        return raw, 1
    low, high = min(numbers), max(numbers)
    if low == high:
        return str(low), 1
    return f"{low}-{high}", high - low + 1


def _normalize_weekday(raw: str) -> str:
    text = raw.replace("星期", "").replace("礼拜", "").replace("周", "").strip()
    if text in ("7", "天"):
        return "日"
    if text.isdigit() and 1 <= int(text) <= 7:
        return WEEKDAY_CHARS[int(text) - 1]
    return text


def _parse_date(text: str) -> date | None:
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y年%m月%d日", "%m-%d", "%m/%d"):
        try:
            parsed = datetime.strptime(text, fmt)
        except ValueError:
            continue
        return parsed.date() if "%Y" in fmt else None
    return None


def _match_columns(headers: list[str]) -> dict[str, ColumnMatch]:
    matches: dict[str, ColumnMatch] = {}
    taken: set[int] = set()

    # Exact alias equality first, so an unambiguous header never loses to a
    # substring match somewhere else in the row.
    for index, header in enumerate(headers):
        if not header or any(hint in header for hint in EXCLUDE_HINTS):
            continue
        best: tuple[int, str] | None = None
        for name in FIELD_ORDER:
            for alias in FIELD_ALIASES[name]:
                if header == alias and (best is None or len(alias) > best[0]):
                    best = (len(alias), name)
        if best is not None and best[1] not in matches:
            matches[best[1]] = ColumnMatch(best[1], index, header, "exact")
            taken.add(index)

    # Then containment, preferring the longest alias and the closest header.
    for name in FIELD_ORDER:
        if name in matches:
            continue
        candidates: list[tuple[int, int, int, str]] = []
        for index, header in enumerate(headers):
            if index in taken or not header:
                continue
            if any(hint in header for hint in EXCLUDE_HINTS):
                continue
            for alias in FIELD_ALIASES[name]:
                if alias in header:
                    candidates.append((-len(alias), len(header), index, header))
                    break
        if candidates:
            candidates.sort()
            _, _, index, header = candidates[0]
            matches[name] = ColumnMatch(name, index, header, "alias")
            taken.add(index)

    return matches


def _find_header_row(rows: list[tuple]) -> tuple[int, list[str], dict[str, ColumnMatch]]:
    best_index = -1
    best_headers: list[str] = []
    best_matches: dict[str, ColumnMatch] = {}
    for index, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        headers = [_norm_header(value) for value in row]
        matches = _match_columns(headers)
        if len(matches) > len(best_matches):
            best_index, best_headers, best_matches = index, headers, matches
    if len(best_matches) < MIN_HEADER_MATCHES:
        detected = [item for item in best_headers if item]
        raise ScheduleParseError(
            "未能在前 15 行中找到课表表头，请确认文件是教务系统导出的课表，或手动指定表头行与列对应关系。",
            detected,
            list(FIELD_LABELS),
        )
    return best_index, best_headers, best_matches


def _from_registrar_matrix(
    path: Path | str,
    course_name: str | None,
    teaching_class: str | None,
) -> ScheduleAnalysis:
    """The registrar's .xls is a weekday matrix, so column mapping never applies."""
    from app.services.registrar_schedule import RegistrarParseError, parse_registrar_xls

    try:
        parsed = parse_registrar_xls(path, course_name=course_name, teaching_class=teaching_class)
    except RegistrarParseError as exc:
        raise ScheduleParseError(str(exc), [], []) from exc

    # A teacher's timetable holds every course they teach, so the raw list also
    # names teaching classes belonging to other courses -- offering those would
    # only produce an empty schedule.
    relevant = [
        item for item in parsed.teaching_classes
        if not course_name or item.startswith(course_name)
    ]
    warnings: list[str] = []
    if not teaching_class and len(relevant) > 1:
        warnings.append(
            "这门课有多个教学班：" + "、".join(relevant) + "，当前把它们合并在一起，请选择一个教学班。"
        )
    warnings.append(f"日期由开学日期 {parsed.term_start} 按周次推算，请在预览中核对。")
    return ScheduleAnalysis(
        sessions=[
            ScheduleSession(
                week_no=item.week_no,
                date_text=item.date_text,
                weekday=item.weekday,
                periods=item.periods,
                course_name=item.course_name,
                class_name=item.class_name,
                location=item.location,
                hours=item.hours,
            )
            for item in parsed.sessions
        ],
        header_row_index=0,
        matches=[],
        detected_headers=[],
        unmapped_headers=[],
        warnings=warnings,
        course_names=parsed.course_names,
        teaching_classes=relevant,
    )


def analyze_schedule(
    path: Path | str,
    *,
    header_row: int | None = None,
    mapping: dict[str, int] | None = None,
    course_name: str | None = None,
    teaching_class: str | None = None,
) -> ScheduleAnalysis:
    if Path(path).suffix.lower() == ".xls":
        return _from_registrar_matrix(path, course_name, teaching_class)

    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ScheduleParseError("课表文件为空", [], list(FIELD_LABELS))

    if mapping is not None:
        header_index = header_row if header_row is not None else 0
        headers = [_norm_header(value) for value in rows[header_index]]
        matches = {
            name: ColumnMatch(name, index, headers[index] if index < len(headers) else "", "manual")
            for name, index in mapping.items()
            if name in FIELD_LABELS
        }
    elif header_row is not None:
        headers = [_norm_header(value) for value in rows[header_row]]
        matches = _match_columns(headers)
        header_index = header_row
    else:
        header_index, headers, matches = _find_header_row(rows)

    detected_headers = [item for item in headers if item]
    missing_required = [name for name in REQUIRED_FIELDS if name not in matches]
    if "date_text" not in matches and "week_no" not in matches:
        missing_required.append("date_text")
    if missing_required:
        labels = "、".join(FIELD_LABELS[name] for name in missing_required)
        raise ScheduleParseError(
            f"课表缺少必需的列：{labels}。已识别到的列：{'、'.join(detected_headers) or '（无）'}。",
            detected_headers,
            missing_required,
        )

    warnings: list[str] = []
    fill_fields = [name for name in FIELD_ORDER if name != "periods" and name in matches]
    carried: dict[str, str] = {}
    raw_rows: list[dict[str, str]] = []

    for row in rows[header_index + 1 :]:
        values: dict[str, str] = {}
        for name in fill_fields:
            index = matches[name].column_index
            text = _norm(row[index]) if index < len(row) else ""
            if text:
                carried[name] = text
            values[name] = text or carried.get(name, "")
        periods_index = matches["periods"].column_index
        periods_raw = _norm(row[periods_index]) if periods_index < len(row) else ""
        if not periods_raw:
            continue
        values["periods"] = periods_raw
        raw_rows.append(values)

    course_names: list[str] = []
    for values in raw_rows:
        name = values.get("course_name", "")
        if name and name not in course_names:
            course_names.append(name)

    if course_name:
        wanted = _norm(course_name)
        raw_rows = [
            values
            for values in raw_rows
            if not values.get("course_name")
            or values["course_name"] == wanted
            or wanted in values["course_name"]
            or values["course_name"] in wanted
        ]

    if "weekday" not in matches:
        warnings.append("课表没有星期列，已根据日期推算星期，请在预览中核对。")
    if "week_no" not in matches:
        warnings.append("课表没有周次列，已按上课日期顺序编号，请在预览中核对。")

    week_by_date: dict[str, int] = {}
    if "week_no" not in matches:
        for values in raw_rows:
            key = values.get("date_text", "")
            if key and key not in week_by_date:
                week_by_date[key] = len(week_by_date) + 1

    sessions: list[ScheduleSession] = []
    for position, values in enumerate(raw_rows, start=1):
        date_text = values.get("date_text", "")
        periods_text, hours = _normalize_periods(values["periods"])

        weekday = _normalize_weekday(values.get("weekday", ""))
        if not weekday:
            parsed_date = _parse_date(date_text)
            if parsed_date is not None:
                weekday = WEEKDAY_CHARS[parsed_date.weekday()]

        raw_week = values.get("week_no", "")
        digits = re.findall(r"\d+", raw_week)
        if digits:
            week_no = int(digits[0])
        elif date_text and date_text in week_by_date:
            week_no = week_by_date[date_text]
        else:
            week_no = position

        sessions.append(
            ScheduleSession(
                week_no=week_no,
                date_text=date_text,
                weekday=weekday,
                periods=periods_text,
                course_name=values.get("course_name", ""),
                class_name=values.get("class_name", ""),
                location=values.get("location", ""),
                hours=hours,
            )
        )

    ordered_matches = [matches[name] for name in FIELD_ORDER if name in matches]
    taken = {match.column_index for match in ordered_matches}
    unmapped = [header for index, header in enumerate(headers) if header and index not in taken]

    return ScheduleAnalysis(
        sessions=sessions,
        header_row_index=header_index,
        matches=ordered_matches,
        detected_headers=detected_headers,
        unmapped_headers=unmapped,
        warnings=warnings,
        course_names=course_names,
    )


def parse_schedule(
    path: Path | str,
    *,
    header_row: int | None = None,
    mapping: dict[str, int] | None = None,
    course_name: str | None = None,
    teaching_class: str | None = None,
) -> list[ScheduleSession]:
    return analyze_schedule(
        path,
        header_row=header_row,
        mapping=mapping,
        course_name=course_name,
        teaching_class=teaching_class,
    ).sessions
