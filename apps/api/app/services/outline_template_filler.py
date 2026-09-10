"""Write the generated body of the outline into the school's template.

学习进程 used to be the only generated part; every other section shipped as
whatever the uploaded template happened to say. The school's standard template
marks each spot it expects filled with 【AI 生成…】 in the same dark red the
lesson template uses, so this walks those markers, writes what belongs there,
repeats a line or a table row where there are several items, and drops the red
-- the colour is the school telling us what to fill, not part of the document it
wants back.

Not everything written here is generated: 五、学习资源 and 六、考核方式 are copied
out of the course standard word for word. See `course_standard_resources` and
`course_standard_assessment` for why.

A template carrying no marker -- a teacher's own finished outline -- is left
exactly as it was before: nothing here matches it, so nothing here writes to it.
"""

from collections.abc import Sequence
from copy import deepcopy
from dataclasses import dataclass
import re
from typing import Any

from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.services.course_standard_assessment import CourseAssessments, assessment_rows
from app.services.course_standard_resources import MISSING, CourseResources, resource_lines
from app.services.lesson_template_filler import (
    GENERATED_MARK_COLOR,
    dominant_run_properties,
    set_cell_text,
)


GENERATED_PLACEHOLDER = re.compile(r"【\s*AI\s*生成[^】]*】")
# 「[1] 」 and 「1．」 -- an entry the standard already numbers is not numbered again.
LEADING_NUMBER = re.compile(r"^\s*(\[\d+\]|\d{1,2}\s*[.．、])")
SECTION_HEADING = re.compile(r"^\s*[一二三四五六七八九十]+\s*[、.．]")
SUBSECTION_HEADING = re.compile(r"^\s*[（(]\s*[一二三四五六七八九十]+\s*[）)]")


class OutlineTemplateError(ValueError):
    pass


@dataclass(frozen=True)
class _ParagraphSlot:
    """Where one generated paragraph goes.

    ``label`` is text the template writes ahead of the marker, either on the
    same line (「课程简介：」) or as the line above it (「教材：」 heading a list);
    ``heading`` is the 一、/（一） heading in force above it. Sections that carry
    a label are told apart by it, the rest by their heading. ``numbered`` lists
    are written 「[1] 」「[2] 」 as the school's finished outlines cite them.
    """

    field: str
    label: str = ""
    heading: str = ""
    repeat: bool = False
    numbered: bool = False


# In template order. 五、学习资源 is copied from the course standard rather than
# written, so its four subsections name themselves as `resource_lines()` does.
PARAGRAPH_SLOTS: tuple[_ParagraphSlot, ...] = (
    _ParagraphSlot("course_summary", label="课程简介"),
    _ParagraphSlot("teaching_strategy", label="教学策略"),
    _ParagraphSlot("prerequisites", label="先修要求"),
    _ParagraphSlot("learning_outcomes", heading="预期学习成果", repeat=True),
    _ParagraphSlot("教材", label="教材", repeat=True, numbered=True),
    _ParagraphSlot("教辅", label="教辅", repeat=True, numbered=True),
    _ParagraphSlot("在线学习资源", heading="在线学习资源", repeat=True),
    _ParagraphSlot("高质量作业范例", heading="作业范例", repeat=True),
    _ParagraphSlot("参考书目", heading="参考书目", repeat=True, numbered=True),
    _ParagraphSlot("study_advice", heading="学习方法"),
    _ParagraphSlot("academic_integrity", heading="学术诚信"),
    _ParagraphSlot("attendance", heading="出勤"),
    _ParagraphSlot("classroom_discipline", heading="课堂纪律"),
    _ParagraphSlot("assignment_requirements", heading="学习任务"),
)

CONTENT_TABLE_COLUMNS = (
    ("教学单元", "name"),
    ("参考课时", "hours"),
    ("重点", "key_points"),
    ("难点", "key_points"),
    ("教学方法", "teaching_methods"),
    ("教学内容", "teaching_content"),
)
CONTENT_TABLE_REQUIRED = {"name", "hours"}
ASSESSMENT_TABLE_MARKERS = ("考核项目", "比例")


@dataclass(frozen=True)
class _TableRow:
    """One row to write: its text by grid column, and how its cells join up."""

    cells: dict[int, str]
    spans: tuple[tuple[int, ...], ...] = ()
    continues_above: tuple[int, ...] = ()


LABEL_LINE = re.compile(r"^\s*(?P<label>[^：:]{2,12}?)\s*[：:]\s*(?P<value>.*)$")


def fill_label_lines(document, values: dict[str, str]) -> list[str]:
    """Write 「教师姓名：」-style lines the template leaves blank after the colon.

    The 课程信息 / 教师信息 block carries no marker: the school's template
    simply prints the label and expects the teacher to type after it. A line
    that already has something after the colon is the teacher's own and is
    left alone, so a finished outline used as a template keeps its text.
    Returns the labels written.
    """
    written: list[str] = []
    for paragraph in document.paragraphs:
        match = LABEL_LINE.match(paragraph.text)
        if match is None:
            continue
        label = "".join(match.group("label").split())
        value = values.get(label, "")
        if not value.strip() or match.group("value").strip():
            continue
        _append_after_label(paragraph, value)
        written.append(label)
    return written


def _append_after_label(paragraph: Paragraph, value: str) -> None:
    # The label's own run carries the font the school chose for the line, so
    # the value continues in that run instead of starting a default-styled one.
    run = paragraph.runs[-1] if paragraph.runs else paragraph.add_run()
    first, *rest = value.split("\n")
    run.text = run.text + first
    for line in rest:
        run.add_break()
        run.add_text(line)


def template_marks_generated_sections(document) -> bool:
    """Whether this template asks for a generated body at all.

    The caller checks before spending a model call: a teacher's own outline
    carries no marker and needs nothing generated beyond 学习进程.
    """
    return bool(find_unfilled_generated_sections(document))


def fill_outline_sections(
    document,
    sections: Any,
    resources: CourseResources | None = None,
    assessments: CourseAssessments | None = None,
) -> bool:
    """Write every marked section. Returns whether anything was written."""
    values = _section_values(sections, resources)
    if not values:
        return False

    written = False
    used: set[int] = set()
    for paragraph, heading, lead_in in _marked_paragraphs(document):
        marker = GENERATED_PLACEHOLDER.search(paragraph.text)
        if marker is None:
            continue
        matched = _match_slot(heading, paragraph.text[: marker.start()], lead_in, used)
        if matched is None:
            continue
        index, slot = matched
        value = values.get(slot.field)
        if value is None:
            continue
        used.add(index)
        _write_slot(paragraph, slot, value)
        written = True

    written |= _fill_content_table(document, sections.learning_units)
    written |= _fill_assessment_table(document, assessments)
    return written


def find_unfilled_generated_sections(document) -> list[str]:
    """Markers the export left standing.

    The same bargain the lesson template strikes: the school states in the
    template which parts it expects generated, so a marker still sitting there
    means a section nobody wired up would ship as an instruction to the teacher.
    """
    stale = [
        (heading or paragraph.text.strip())[:24]
        for paragraph, heading, _lead_in in _marked_paragraphs(document)
    ]
    for table in document.tables:
        for row in table.rows:
            if _row_is_marked(row):
                header = _normalize(table.rows[0].cells[0].text)[:12] or "表格"
                stale.append(f"「{header}」所在表格")
                break
    return list(dict.fromkeys(stale))


def _section_values(sections: Any, resources: CourseResources | None) -> dict[str, Any]:
    if sections is None:
        return {}
    values: dict[str, Any] = {
        "course_summary": sections.course_summary,
        "teaching_strategy": sections.teaching_strategy,
        "prerequisites": sections.prerequisites,
        "learning_outcomes": list(sections.learning_outcomes),
        "study_advice": sections.study_advice,
        "academic_integrity": sections.academic_integrity,
        "attendance": sections.attendance,
        "classroom_discipline": sections.classroom_discipline,
        "assignment_requirements": sections.assignment_requirements,
    }
    # No course standard to copy from is not a reason to invent resources; the
    # extractor's own "please supply this" line is the honest answer.
    values.update(resource_lines(resources if resources is not None else CourseResources()))
    return values


def _unit_values(unit: Any) -> dict[str, str]:
    return {
        "name": str(unit.name or ""),
        "teaching_content": str(unit.teaching_content or ""),
        "teaching_methods": str(unit.teaching_methods or ""),
        "key_points": str(unit.key_points or ""),
        "hours": str(unit.hours or ""),
    }


def _normalize(value: str) -> str:
    return "".join(value.split())


def _marked_paragraphs(document) -> list[tuple[Paragraph, str, str]]:
    """Marked paragraphs, each with the headings in force above it and the
    unmarked line just before it -- 「教材：」 heads its list from its own line."""
    found: list[tuple[Paragraph, str, str]] = []
    section = ""
    subsection = ""
    lead_in = ""
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if GENERATED_PLACEHOLDER.search(text) is None:
            if SECTION_HEADING.match(text):
                section, subsection = text, ""
            elif SUBSECTION_HEADING.match(text):
                subsection = text
            lead_in = text
            continue
        found.append((paragraph, f"{section} {subsection}".strip(), lead_in))
        lead_in = ""
    return found


def _marked_tables(document) -> list[tuple[Table, int]]:
    """Tables holding a marker, paired with the index of their header row."""
    found: list[tuple[Table, int]] = []
    for table in document.tables:
        marked = [index for index, row in enumerate(table.rows) if _row_is_marked(row)]
        if marked and marked[0] > 0:
            found.append((table, marked[0] - 1))
    return found


def _row_is_marked(row) -> bool:
    return any(GENERATED_PLACEHOLDER.search(cell.text) for cell in row.cells)


def _match_slot(heading: str, prefix: str, lead_in: str, used: set[int]) -> tuple[int, _ParagraphSlot] | None:
    for index, slot in enumerate(PARAGRAPH_SLOTS):
        if index in used:
            continue
        if slot.label and slot.label not in prefix and not lead_in.startswith(slot.label):
            continue
        if slot.heading and slot.heading not in heading:
            continue
        return index, slot
    return None


def _write_slot(paragraph: Paragraph, slot: _ParagraphSlot, value: Any) -> None:
    if not slot.repeat:
        _write_placeholder(paragraph, str(value))
        return
    items = [str(item) for item in value if str(item).strip()]
    if not items:
        paragraph._p.getparent().remove(paragraph._p)
        return
    if slot.numbered:
        items = _numbered(items)
    # The template holds one line and says to copy it per item, so the second
    # entry keeps the first one's indent and spacing rather than the default.
    # Its label does not come along: 「教辅：」 heads the list once, as it does in
    # the finished outlines this template was cleaned from.
    clones = _repeat_paragraph(paragraph, len(items))
    for position, (written, clone) in enumerate(zip(items, clones)):
        _write_placeholder(clone, written, from_line_start=position > 0)


def _numbered(items: list[str]) -> list[str]:
    """「[1] 」 ahead of each citation, the way the finished outlines cite.

    The standard's own numbering, where it has one, is kept as it is; and the
    line asking the teacher to supply the entry is a note, not a citation.
    """
    if any(LEADING_NUMBER.match(item) for item in items) or items == [MISSING]:
        return items
    return [f"[{position}] {item.strip()}" for position, item in enumerate(items, start=1)]


def _repeat_paragraph(paragraph: Paragraph, count: int) -> list[Paragraph]:
    clones = [paragraph]
    anchor = paragraph._p
    for _ in range(max(count, 1) - 1):
        copied = deepcopy(paragraph._p)
        anchor.addnext(copied)
        anchor = copied
        clones.append(Paragraph(copied, paragraph._parent))
    return clones


def _write_placeholder(paragraph: Paragraph, text: str, from_line_start: bool = False) -> None:
    """Replace the marker in place, keeping whatever the line says around it."""
    marker = GENERATED_PLACEHOLDER.search(paragraph.text)
    if marker is None:
        return
    start, end = (0 if from_line_start else marker.start()), marker.end()
    offset = 0
    replaced = False
    for run in paragraph.runs:
        run_start, run_end = offset, offset + len(run.text)
        offset = run_end
        if run_end <= start or run_start >= end:
            continue
        head = run.text[: max(0, start - run_start)]
        tail = run.text[max(0, end - run_start) :] if run_end > end else ""
        run.text = head + text + tail if not replaced else head + tail
        replaced = True
        _strip_mark_colour(run)


def _strip_mark_colour(run) -> None:
    """The red said "generate this"; the finished document is black."""
    properties = run._r.find(qn("w:rPr"))
    if properties is None:
        return
    colour = properties.find(qn("w:color"))
    if colour is not None and colour.get(qn("w:val")) == GENERATED_MARK_COLOR:
        properties.remove(colour)


def clear_generated_marking(document) -> None:
    """Drop the contract colour from anything the template still shows in it.

    The template paints whole lines red, lead-in sentences and paragraph marks
    included -- 「通过本课程学习，同学们能够：」 carries no marker but is red all
    the same, and a red paragraph mark turns the teacher's own next sentence
    red. Only the exact contract colour is touched, so a template that colours
    something of its own keeps it.
    """
    for colour in list(document.element.body.iter(qn("w:color"))):
        if colour.get(qn("w:val")) == GENERATED_MARK_COLOR:
            colour.getparent().remove(colour)


def _fill_content_table(document, units: Sequence[Any]) -> bool:
    for table, header_index in _marked_tables(document):
        columns = _columns_by_header(table.rows[header_index], CONTENT_TABLE_COLUMNS)
        if not CONTENT_TABLE_REQUIRED.issubset(columns.values()):
            continue
        _write_table_rows(
            table,
            header_index,
            [
                _TableRow(cells={index: _unit_values(unit)[field] for index, field in columns.items()})
                for unit in units
            ],
        )
        return True
    return False


def _fill_assessment_table(document, assessments: CourseAssessments | None) -> bool:
    """六、考核方式 is the course standard's own table, copied across.

    Text, column merges and the 合计 row all come from the standard: the school
    decided this split of the marks, and it is the split the 教务 record checks.
    """
    for table, header_index in _marked_tables(document):
        header = _normalize("".join(cell.text for cell in table.rows[header_index].cells))
        if not all(marker in header for marker in ASSESSMENT_TABLE_MARKERS):
            continue
        width = len(table.rows[header_index].cells)
        rows = assessment_rows(assessments or CourseAssessments(), width)
        stated = max((len(row.cells) for row in rows), default=0)
        if stated > width:
            # 比例 is the standard's last column. Writing what fits would drop it
            # and still look like a finished table.
            raise OutlineTemplateError(
                f"导出中止：课程标准的考核方式表有 {stated} 列，模板的只有 {width} 列，"
                "照抄会丢掉最后一列（比例）。请改用与课程标准同一套的大纲模板。"
            )
        _write_table_rows(
            table,
            header_index,
            [
                _TableRow(
                    cells=dict(enumerate(row.cells)),
                    spans=row.spans,
                    continues_above=row.continues_above,
                )
                for row in rows
            ],
        )
        return True
    return False


def _columns_by_header(row, headers: Sequence[tuple[str, str]]) -> dict[int, str]:
    columns: dict[int, str] = {}
    for index, cell in enumerate(row.cells):
        text = _normalize(cell.text)
        for keyword, field in headers:
            if keyword in text:
                columns[index] = field
                break
    return columns


def _write_table_rows(table: Table, header_index: int, rows: Sequence[_TableRow]) -> None:
    """Grow the marked row to one row per item, then write them.

    Only rows that carry a marker are cloned or dropped, so a 合计 row the
    template draws under the data survives untouched.
    """
    elements = [row._tr for row in table.rows[header_index + 1 :] if _row_is_marked(row)]
    if not elements:
        return
    template_row = elements[-1]
    anchor = template_row
    while len(elements) < len(rows):
        copied = deepcopy(template_row)
        anchor.addnext(copied)
        anchor = copied
        elements.append(copied)
    for surplus in elements[len(rows) :]:
        surplus.getparent().remove(surplus)
    del elements[len(rows) :]

    # Merge before writing: merging concatenates what the cells hold, and every
    # clone still holds the template's 【AI 生成】 at this point.
    for element in elements:
        _clear_vertical_merges(element)
    by_element = {row._tr: row for row in table.rows}
    positions = {row._tr: index for index, row in enumerate(table.rows)}
    for element, plan in zip(elements, rows):
        _apply_spans(by_element[element], plan.spans)
    _merge_down(table, [positions[element] for element in elements], rows)

    donor = dominant_run_properties(table)
    for element, plan in zip(elements, rows):
        row = by_element[element]
        for index, text in plan.cells.items():
            # Re-read the cells each time: a merge changed how many there are,
            # and a merged cell answers to every column it covers.
            cells = row.cells
            if index < len(cells):
                # Each marked cell states its own format -- 教学单元 a size up
                # from the columns beside it -- so a cell that carries one keeps
                # it; the table's donor is for a blank cell only.
                set_cell_text(cells[index], text, None if _has_run_properties(cells[index]) else donor)


def _has_run_properties(cell) -> bool:
    paragraph = cell.paragraphs[0]
    return bool(paragraph.runs) and paragraph.runs[0]._r.find(qn("w:rPr")) is not None


def _clear_vertical_merges(element) -> None:
    """Drop the template's own w:vMerge before laying down the real one.

    The school's blank template marks its single data row as the start of a
    vertical merge; cloned twenty times that leaves twenty stray group starts.
    """
    for cell in element.findall(qn("w:tc")):
        properties = cell.find(qn("w:tcPr"))
        merge = None if properties is None else properties.find(qn("w:vMerge"))
        if merge is not None:
            properties.remove(merge)


def _apply_spans(row, spans: Sequence[Sequence[int]]) -> None:
    for span in spans:
        if len(span) < 2:
            continue
        merged = row.cells[span[0]]
        for index in span[1:]:
            candidate = row.cells[index]
            if candidate._tc is not merged._tc:
                merged = merged.merge(candidate)


def _merge_down(table: Table, row_indexes: Sequence[int], rows: Sequence[_TableRow]) -> None:
    for column, first, last in _vertical_spans(rows):
        table.cell(row_indexes[first], column).merge(table.cell(row_indexes[last], column))


def _vertical_spans(rows: Sequence[_TableRow]) -> list[tuple[int, int, int]]:
    """Runs of rows sharing one cell in a column, as (column, first, last)."""
    found: list[tuple[int, int, int]] = []
    for column in sorted({column for row in rows for column in row.continues_above}):
        position = 0
        while position < len(rows):
            if column not in rows[position].continues_above:
                position += 1
                continue
            first = position - 1
            while position < len(rows) and column in rows[position].continues_above:
                position += 1
            # A run continuing from the header has nothing here to merge into.
            if first >= 0:
                found.append((column, first, position - 1))
    return found
